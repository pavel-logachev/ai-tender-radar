"""The manager receives an actionable, safe Excel report, without office dependencies."""
import copy
import io
import unittest
import zipfile
from datetime import datetime, timezone
from xml.etree import ElementTree as ET

NOW = datetime(2026, 9, 30, 6, 5, tzinfo=timezone.utc)
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def candidate(key="ready", *, contacts=True, status="proposal"):
    return {"card": {"id": key, "title": "Поставка серверов", "customer_name": "Завод",
                     "source_url": "https://tenders.example.org/process/" + key},
            "fingerprint": key, "analysis_state": "complete", "texts": 1, "gaps": 0,
            "analysis": {"interesting": True, "summary": "Расширение инфраструктуры",
                         "evidence": [{"source": "doc:1", "quote": "Система хранения данных 24 ТБ"}],
                         "document_findings": ["Система хранения данных 24 ТБ"],
                         "questions": ["Нужна ли миграция?"]},
            "opportunity": {"status": status, "acceptance_end_date": "2026-10-02T12:00:00Z",
                "buyer": {"inn": "012345678901", "legal_name": "Завод"},
                "budget": {"decision": "unknown_include"},
                "customer_intelligence": {"contacts": [{"phone": "+7 (900) 000-33-55",
                    "email": "purchase@example.org", "name": None, "role": None,
                    "source_url": "https://example.org/contacts", "source_kind": "official_site"}] if contacts else [],
                    "finance": {"state": "verified", "year": 2025, "revenue_rub": "100000000",
                        "net_profit_rub": "0", "equity_rub": None,
                        "source_url": "https://bo.nalog.gov.ru/organizations-card/123"}, "gaps": []}}}


def prepared(cards):
    return {"slot": "2026-09-30-09", "collected_at": "2026-09-30T06:00:00Z", "cards": cards,
            "stale": False, "considered": len(cards), "history_total": len(cards),
            "analysis_failures": 0, "coverage": {"response_rows": 15000, "termination": "short_page"}}


def sheets(payload):
    with zipfile.ZipFile(io.BytesIO(payload)) as book:
        names = [s.attrib["name"] for s in ET.fromstring(book.read("xl/workbook.xml")).findall("s:sheets/s:sheet", NS)]
        return {name: ET.fromstring(book.read(f"xl/worksheets/sheet{i}.xml")) for i, name in enumerate(names, 1)}


def text(sheet):
    return "\n".join(node.text or "" for node in sheet.findall(".//s:t", NS))


def rows(sheet):
    return [["".join(c.itertext()) for c in r.findall("s:c", NS)] for r in sheet.findall("s:sheetData/s:row", NS) if int(r.attrib["r"]) >= 5]


class ExcelReportTest(unittest.TestCase):
    def test_only_contact_ready_active_hardware_is_a_lead_other_cards_remain_visible(self):
        from agent_radar.excel_report import build_excel_report
        expired = candidate("expired")
        expired["opportunity"]["acceptance_end_date"] = "2026-09-30T06:04:00Z"
        pending = candidate("pending")
        pending.update(analysis=None, analysis_state="pending")
        cards = [candidate(), candidate("no-contact", contacts=False), candidate("ended", status="summarizing"), expired, pending]
        caption, payload, seen = build_excel_report(prepared(cards), now=NOW)
        result = sheets(payload)
        self.assertEqual(list(result), ["Лиды", "Доработка", "История", "Детали", "Сводка"])
        self.assertEqual(len(rows(result["Лиды"])), 1)
        self.assertEqual(len(rows(result["Доработка"])), 2)
        self.assertEqual(len(rows(result["История"])), 2)
        self.assertIn("Нет опубликованного рабочего контакта", text(result["Доработка"]))
        self.assertIn("Разбор не завершён", text(result["Доработка"]))
        self.assertIn("Готовых к первому обращению: 1", caption)
        self.assertEqual(seen, [c["fingerprint"] for c in cards])
        self.assertNotIn("от5", caption)
        self.assertIn("Бюджет неизвестен", text(result["Лиды"]))
        self.assertIn("14 дней", text(result["Сводка"]))
        self.assertIn("5 минут", text(result["Сводка"]))
        self.assertIn("lastModified", text(result["Сводка"]))
        self.assertNotIn('"customer_intelligence"', text(result["Лиды"]))

    def test_xlsx_is_styled_filtered_numeric_and_safe_for_untrusted_source_text(self):
        from agent_radar.excel_report import build_excel_report
        c = candidate()
        c["card"]["customer_name"] = '=HYPERLINK("file:///secret")\x00<|im_end|>'
        c["opportunity"]["buyer"]["legal_name"] = c["card"]["customer_name"]
        c["analysis"]["summary"] = "@SUM(A1:A2)"
        c["opportunity"]["budget"] = {"amount": "6000000", "currency": "RUB", "basis": "procedure_total"}
        _, payload, _ = build_excel_report(prepared([c]), now=NOW)
        result = sheets(payload)
        main = result["Лиды"]
        cell = main.find(".//s:c[@r='D5']", NS)
        self.assertIsNone(cell.get("t"))
        self.assertAlmostEqual(float(cell.find("s:v", NS).text), 46297.625)
        budget_row = next(r for r in result["Детали"].findall("s:sheetData/s:row", NS) if "Бюджет закупки" in "".join(r.itertext()))
        self.assertEqual(budget_row.findall("s:c", NS)[3].find("s:v", NS).text, "6000000")
        self.assertIn("6 000 000 ₽", text(main))
        self.assertEqual(main.find(".//s:pane", NS).get("topLeftCell"), "C5")
        self.assertEqual(main.find("s:autoFilter", NS).get("ref"), "A4:J5")
        self.assertEqual(main.find("s:pageSetup", NS).get("paperSize"), "8")
        self.assertEqual(main.find("s:pageSetup", NS).get("fitToWidth"), "1")
        self.assertEqual(main.find("s:pageSetup", NS).get("orientation"), "landscape")
        self.assertGreaterEqual(float(main.find("s:sheetData/s:row[@r='5']", NS).get("ht")), 100)
        self.assertEqual(main.find(".//s:c[@r='B5']", NS).get("t"), "inlineStr")
        self.assertIn('=HYPERLINK("file:///secret")', text(main))
        self.assertNotIn("<|im_end|>", text(main))
        self.assertIsNone(main.find(".//s:f", NS))
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            self.assertLessEqual(len(archive.namelist()), 200)
            self.assertLessEqual(sum(z.file_size for z in archive.infolist()), 10_000_000)
            self.assertLessEqual(len(payload), 2_000_000)
            self.assertFalse(any(x.endswith(".bin") or "embedding" in x or "externalLink" in x for x in archive.namelist()))
            table = ET.fromstring(archive.read("xl/tables/table1.xml"))
            self.assertEqual(table.get("ref"), "A4:J5")
            self.assertEqual(table.find("s:tableStyleInfo", NS).get("showRowStripes"), "1")
            styles = ET.fromstring(archive.read("xl/styles.xml"))
            self.assertTrue(styles.findall("s:numFmts/s:numFmt", NS))
            links = ET.fromstring(archive.read("xl/worksheets/_rels/sheet1.xml.rels"))
            external = [x.get("Target") for x in links if x.get("TargetMode") == "External"]
            self.assertIn("https://tenders.example.org/process/ready", external)
            self.assertTrue(all(x.startswith("https://") for x in external))
        finance = [r for r in rows(result["Детали"]) if r[2] == "Чистая прибыль, RUB"][0]
        self.assertEqual(finance[3], "0")
        self.assertEqual(finance[5], "2025")
        self.assertIn("012345678901", text(result["Детали"]))
        self.assertIn("Неизвестно", text(result["Детали"]))

    def test_stale_keeps_all_active_candidates_in_work_and_fingerprints_history(self):
        from agent_radar.excel_report import build_excel_report
        p = prepared([candidate(), candidate("waiting", contacts=False), candidate("history", status="ended")])
        caption, payload, seen = build_excel_report(p, now=datetime(2026, 9, 30, 7, 0, tzinfo=timezone.utc))
        result = sheets(payload)
        self.assertIn("Готовых к первому обращению: 0", caption)
        self.assertIn("Готовых лидов нет", text(result["Лиды"]))
        self.assertEqual(len(rows(result["Доработка"])), 2)
        self.assertEqual(len(rows(result["История"])), 1)
        self.assertIn("Данные устарели", text(result["Доработка"]))
        self.assertEqual(seen, ["ready", "waiting", "history"])

    def test_contact_provenance_channels_status_profile_and_deadline_are_conservative(self):
        from agent_radar.excel_report import build_excel_report
        bad_urls = ["https://user:pass@example.org/contacts", "https://example.org/contact?token=secret",
                    "https://127.0.0.1/contact", "https://x.internal/contact", "http://example.org/contact",
                    "https://example.org/contact\x00", "https://2130706433/contact", "https://example.org/contact#secret"]
        cards = []
        for n, url in enumerate(bad_urls):
            c = candidate("unsafe" + str(n))
            c["opportunity"]["customer_intelligence"]["contacts"][0]["source_url"] = url
            cards.append(c)
        no_channel = candidate("fake")
        no_channel["opportunity"]["customer_intelligence"]["contacts"][0].update(phone="000-00", email="not email", name="Предполагаемый директор")
        cards.append(no_channel)
        c = candidate("guess-kind")
        c["opportunity"]["customer_intelligence"]["contacts"][0]["source_kind"] = "guessed"
        cards.append(c)
        for key, change in [("unknown-status", {"status": None}), ("bad-date", {"acceptance_end_date": "tomorrow"})]:
            c = candidate(key)
            c["opportunity"].update(change)
            cards.append(c)
        service = candidate("software")
        service["card"]["title"] = "Лицензии Windows Server"
        cards.append(service)
        rejected = candidate("not-interesting")
        rejected["analysis"]["interesting"] = False
        cards.append(rejected)
        no_deadline = candidate("no-deadline")
        no_deadline["opportunity"].pop("acceptance_end_date")
        cards.append(no_deadline)
        _, payload, seen = build_excel_report(prepared(cards), now=NOW)
        result = sheets(payload)
        self.assertEqual(len(rows(result["Лиды"])), 1)
        self.assertEqual(rows(result["Лиды"])[0][0], "no-deadline")
        self.assertEqual(len(rows(result["Доработка"])), 14)
        self.assertEqual(len(seen), 15)
        self.assertNotIn("Предполагаемый директор", text(result["Детали"]))

    def test_budgets_and_full_date_day_keep_unknown_separate_from_zero_finance(self):
        from agent_radar.excel_report import build_excel_report
        cheap = candidate("cheap")
        cheap["opportunity"]["budget"] = {"amount": "4999999", "currency": "RUB", "basis": "procedure_total"}
        day = candidate("whole-day")
        day["opportunity"]["acceptance_end_date"] = "2026-09-30"
        day["analysis"]["document_findings"].append("5000 GPU not in evidence")
        _, payload, _ = build_excel_report(prepared([cheap, day]), now=NOW)
        result = sheets(payload)
        self.assertEqual(len(rows(result["Лиды"])), 1)
        self.assertIn("ниже порога", text(result["Доработка"]))
        self.assertNotIn("5000 GPU", text(result["Детали"]))
        self.assertNotIn("5000 GPU", text(result["Лиды"]))
        self.assertAlmostEqual(float(result["Лиды"].find(".//s:c[@r='D5']/s:v", NS).text), 46295.999988426, places=6)
        with self.assertRaises(ValueError):
            build_excel_report(prepared([]), now=NOW.replace(tzinfo=None))

    def test_hundred_cards_are_bounded_without_mutation_and_formula_cells(self):
        from agent_radar.excel_report import build_excel_report
        p = prepared([candidate(str(n)) for n in range(101)])
        original = copy.deepcopy(p)
        _, payload, seen = build_excel_report(p, now=NOW)
        self.assertEqual(p, original)
        self.assertEqual(len(seen), 100)
        self.assertEqual(len(rows(sheets(payload)["Лиды"])), 100)
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            self.assertLessEqual(sum(z.file_size for z in archive.infolist()), 10_000_000)
            self.assertLessEqual(len(payload), 2_000_000)
            for name in archive.namelist():
                if name.endswith(".xml"):
                    root = ET.fromstring(archive.read(name))
                    self.assertIsNone(root.find(".//s:f", NS))
                    self.assertIsNone(root.find(".//s:definedName[@name='HYPERLINK']", NS))

    def test_internal_tokens_are_removed_and_published_extension_phone_remains_actionable(self):
        from agent_radar.excel_report import build_excel_report
        c = candidate()
        c["opportunity"]["customer_intelligence"]["contacts"][0].update(phone="8 (900) 000-11-33, доб. 702-349", email=None)
        bank = candidate("bank-extension")
        bank["opportunity"]["customer_intelligence"]["contacts"][0].update(phone="+7(900)000-44-66 (вн.2831)", email=None)
        self.assertIn("Готовых к первому обращению: 1", build_excel_report(prepared([bank]), now=NOW)[0])
        c["analysis"]["summary"] = "<reasoning>текст</reasoning> [INST] [analysis] <<SYS>> <tool_call> [final]"
        c["card"]["title"] += "\x01 <scratchpad> <|unterminated"
        caption, payload, seen = build_excel_report(prepared([c]), now=NOW)
        result = sheets(payload)
        self.assertIn("Готовых к первому обращению: 1", caption)
        self.assertIn("702-349", text(result["Лиды"]))
        for sheet in result.values():
            for marker in ("<reasoning", "[INST]", "[analysis]", "<<SYS>>", "<tool_call", "[final]", "<scratchpad", "<|"):
                self.assertNotIn(marker, text(sheet))
        self.assertEqual(seen, ["ready"])

    def test_summary_explains_coverage_limit_gaps_and_finance_missing_without_deleting_cards(self):
        from agent_radar.excel_report import build_excel_report
        c = candidate()
        c["opportunity"]["customer_intelligence"]["finance"] = {"state": "unavailable"}
        c["opportunity"]["customer_intelligence"]["contacts"][0].update(phone="0000000000", email=None)
        p = prepared([c])
        p.update(history_total=250, considered=100, analysis_failures=10, source_gap_before_24h=True)
        caption, payload, seen = build_excel_report(p, now=NOW)
        result = sheets(payload)
        self.assertIn("Готовых к первому обращению: 0", caption)
        self.assertIn("Финансовые данные неизвестны", text(result["Сводка"]))
        self.assertIn("150", text(result["Сводка"]))
        self.assertIn("полнота накопления не гарантируется", text(result["Сводка"]))
        self.assertIn("short_page", text(result["Сводка"]))
        self.assertEqual(seen, ["ready"])

    def test_published_contact_source_period_and_financial_risk_remain_explicit(self):
        from agent_radar.excel_report import build_excel_report
        c = candidate()
        contact = c["opportunity"]["customer_intelligence"]["contacts"][0]
        contact.update(source_period="2023", role="Общая приёмная, не технический ЛПР")
        c["opportunity"]["customer_intelligence"]["finance"]["net_profit_rub"] = "-1000"
        _, payload, _ = build_excel_report(prepared([c]), now=NOW)
        result = sheets(payload)
        self.assertIn("Период источника: 2023", text(result["Детали"]))
        self.assertIn("не технический ЛПР", text(result["Лиды"]))
        self.assertIn("Проверить риски", text(result["Лиды"]))
        self.assertIn("не гарантия", text(result["Лиды"]).lower())

    def test_main_table_is_compact_and_general_route_never_claims_a_known_decision_maker(self):
        from agent_radar.excel_report import build_excel_report
        c = candidate()
        c["opportunity"]["customer_intelligence"]["contacts"][0].update(role="Приёмная", route_kind="general_business")
        c["card"]["title"] = "Поставка серверов " + "для аппаратной виртуализации " * 9
        c["opportunity"]["buyer"]["legal_name"] = "ООО Завод " + "промышленного оборудования " * 5
        _, payload, _ = build_excel_report(prepared([c]), now=NOW)
        result = sheets(payload)
        main = result["Лиды"]
        headers = main.find("s:sheetData/s:row[@r='4']", NS).findall("s:c", NS)
        self.assertEqual(len(headers), 10)
        self.assertLessEqual(sum(float(col.get("width")) for col in main.findall("s:cols/s:col", NS)), 235)
        self.assertIn("Контакт / маршрут", text(main))
        self.assertIn("Общий корпоративный маршрут, не закупщик / ИТ-ЛПР", text(main))
        self.assertIn("Выручка 100 млн ₽", text(main))
        self.assertIn("прибыль 0 ₽", text(main))
        self.assertIn("Соедините, пожалуйста", text(result["Детали"]))
        self.assertGreaterEqual(float(main.find("s:sheetData/s:row[@r='5']", NS).get("ht")), 130)

    def test_conflicting_api_and_literal_document_deadlines_need_customer_confirmation(self):
        from agent_radar.excel_report import build_excel_report
        c = candidate("conflict")
        c["opportunity"].update(acceptance_end_date="2026-09-30T14:00:00Z", deadline_conflict={
            "source": "Документация_СХД.docx", "quote": "Приём предложений до 30 сентября 2026 года 12:00 МСК",
            "api_deadline": "2026-09-30T14:00:00Z"})
        caption, payload, seen = build_excel_report(prepared([c]), now=NOW)
        result = sheets(payload)
        self.assertIn("Готовых к первому обращению: 0", caption)
        self.assertEqual(len(rows(result["Доработка"])), 1)
        self.assertEqual(len(rows(result["История"])), 0)
        self.assertIn("Срок в API и документах расходится — уточнить у заказчика", text(result["Доработка"]))
        self.assertIn("Документация_СХД.docx", text(result["Детали"]))
        self.assertIn("Приём предложений до 30 сентября 2026 года 12:00 МСК", text(result["Детали"]))
        self.assertEqual(seen, ["conflict"])

    def test_openpyxl_independent_roundtrip_when_available(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("Optional local Office validator; native runtime has no openpyxl")
        from agent_radar.excel_report import build_excel_report
        _, payload, _ = build_excel_report(prepared([candidate()]), now=NOW)
        book = openpyxl.load_workbook(io.BytesIO(payload))
        self.assertEqual(book.active.title, "Лиды")
        sheet = book["Лиды"]
        self.assertEqual(sheet.freeze_panes, "C5")
        self.assertEqual(sheet.tables["Radar1"].ref, "A4:J5")
        self.assertEqual(sheet["D5"].value, datetime(2026, 10, 2, 15, 0))
        self.assertEqual(sheet["D5"].number_format, "dd.mm.yy hh:mm")
        self.assertEqual(sheet["J5"].hyperlink.target, "https://tenders.example.org/process/ready")
        self.assertEqual(sheet["A5"].hyperlink.location, "'Детали'!A5")
        self.assertEqual(sheet["A1"].fill.fgColor.rgb, "FF112D42")
        self.assertEqual(sheet["A4"].fill.fgColor.rgb, "FF087F8C")
        self.assertEqual(sheet["A5"].font.sz, 11)
        self.assertEqual(sheet["A5"].alignment.wrap_text, True)
        self.assertIsNone(book.vba_archive)
        stream = io.BytesIO()
        book.save(stream)
        reopened = openpyxl.load_workbook(io.BytesIO(stream.getvalue()))
        self.assertEqual(reopened["Лиды"].tables["Radar1"].ref, "A4:J5")
        self.assertFalse(any(c.data_type == "f" for s in reopened for row in s for c in row))


if __name__ == "__main__":
    unittest.main()
