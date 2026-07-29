import importlib
import io
import re
import sys
import types
import unittest
import xml.etree.ElementTree as ET
import zipfile
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


FORBIDDEN_ROLE_WORD = "".join(
    chr(code) for code in (0x0437, 0x0432, 0x043E, 0x043D, 0x0430, 0x0440, 0x044C)
)
FORBIDDEN_COPY_COLUMN = "Карточка для копирования"
INVALID_XML_CONTROL_BYTES_RE = re.compile(rb"[\x00-\x08\x0B\x0C\x0E-\x1F]")
KEY_LEAD_COLUMNS = [
    "ID закупки / external_id",
    "Название",
    "Заказчик",
    "Документы",
    "Сигнал по заказчику",
    "Кого искать",
    "Ссылка",
    "Статус",
    "Примечание менеджера",
]


def make_row(
    tender_id: str,
    *,
    recommendation: str = "maybe",
    llm_report_recommendation: str = "go",
    deadline_at: datetime | None = None,
    docs_count: int = 1,
    docs_with_text: int = 1,
    latest_sales_status: str | None = None,
    latest_sales_status_created_at: datetime | None = None,
    preparation_blocked_event: str | None = None,
    preparation_requeue_event: str | None = None,
    raw_full: dict | None = None,
    lead_triage_decision: str | None = None,
    lead_triage_created_at: datetime | None = None,
    llm_report_created_at: datetime | None = None,
) -> dict:
    row = {
        "tender_id": tender_id,
        "id": tender_id,
        "external_id": tender_id,
        "title": f"Закупка {tender_id}",
        "customer_name": "ООО Заказчик",
        "initial_price": 1250000,
        "region": "Москва",
        "law": "44-ФЗ",
        "procedure_type": "Электронный аукцион",
        "published_at": datetime(2026, 5, 5, 9, 0, tzinfo=timezone.utc),
        "created_at": datetime(2026, 5, 5, 9, 0, tzinfo=timezone.utc),
        "deadline_at": deadline_at,
        "url": "https://example.test/tender",
        "raw": {"full": raw_full or {}},
        "score": 50,
        "recommendation": recommendation,
        "confidence": "high",
        "result": {"why_relevant": ["профильная инфраструктурная закупка"]},
        "rule_based_result": {"why_relevant": ["профильная инфраструктурная закупка"]},
        "document_risk_result": None,
        "llm_report_analysis_type": "llm_customer_lead_report",
        "llm_report_result": {
            "report": {
                "document_status": {
                    "code": "technical_spec_found",
                    "label": "Документы: ТЗ найдено",
                    "analysis_basis": "technical_document",
                },
                "lead_summary": "Заказчик показывает закупочный сигнал по инфраструктуре.",
                "customer_signal": "Идет обновление инфраструктуры.",
                "likely_customer_story": "Вероятно, заказчик обновляет инфраструктурный контур.",
                "possible_needs": ["серверы", "сервисная поддержка", "резервное копирование"],
                "target_end_customer_roles": ["ИТ-директор", "руководитель инфраструктуры"],
                "opening_phrase": "Добрый день. Увидели инфраструктурную закупку и хотели познакомиться с ответственным.",
                "discovery_questions": [
                    "Кто отвечает за инфраструктурное направление?",
                    "Есть ли похожие задачи в плане?",
                    "Кому удобно передать контекст?",
                ],
                "next_action": "Проверить контакт и взять лид в первичную обработку.",
                "lead_priority": "high",
                "confidence": "high",
            }
        },
        "llm_report_recommendation": llm_report_recommendation,
        "llm_report_created_at": llm_report_created_at
        or datetime(2026, 5, 5, 10, 0, tzinfo=timezone.utc),
        "docs_count": docs_count,
        "docs_with_text": docs_with_text,
        "latest_sales_status": latest_sales_status,
        "latest_sales_status_created_at": latest_sales_status_created_at,
        "latest_sales_user_name": "manager" if latest_sales_status else None,
        "latest_sales_comment": "уточнить позже" if latest_sales_status else None,
        "preparation_blocked_event": preparation_blocked_event,
        "preparation_requeue_event": preparation_requeue_event,
    }
    if lead_triage_decision is not None:
        row.update(
            {
                "lead_triage_analysis_type": "llm_customer_lead_triage",
                "lead_triage_recommendation": {
                    "go": "go",
                    "maybe": "maybe",
                    "reject": "no_go",
                }.get(lead_triage_decision, "no_go"),
                "lead_triage_decision": lead_triage_decision,
                "lead_triage_created_at": lead_triage_created_at
                or datetime(2026, 5, 5, 9, 55, tzinfo=timezone.utc),
                "lead_triage_result": {"report": {"lead_decision": lead_triage_decision}},
            }
        )
    return row


class ExcelExportTest(unittest.TestCase):
    def setUp(self) -> None:
        psycopg = types.ModuleType("psycopg")
        psycopg.connect = lambda *args, **kwargs: None

        psycopg_rows = types.ModuleType("psycopg.rows")
        psycopg_rows.dict_row = object()

        config = types.ModuleType("app.config")
        config.settings = SimpleNamespace(database_url="postgresql://user:password@localhost/db")

        business_rules = types.ModuleType("app.business_rules")
        business_rules.business_assessment = lambda tender: {"action": "quick_check"}
        business_rules.business_assessment_block = lambda tender: (
            "Коммерческий интерес: средний\n"
            "Доступность сделки: профильная закупка\n"
            "Эффект заточки: нейтрально\n"
            "Бизнес-действие: быстрый ручной чек перед пресейлом"
        )
        business_rules.digest_sort_key = lambda tender: (0, 0, 0, str(tender.get("title") or ""))
        business_rules.effective_recommendation = lambda tender: tender.get("recommendation") or "go"
        business_rules.lead_hard_noise_reason_for_row = (
            lambda tender, assessment=None: "generic_antivirus_license"
            if "kaspersky" in str(tender.get("title") or "").lower()
            else None
        )
        business_rules.selection_warnings = lambda tender: []

        self.previous_modules = {
            name: sys.modules.get(name)
            for name in ("app.digest", "app.excel_export", "app.lead_notes")
        }
        for name in self.previous_modules:
            sys.modules.pop(name, None)

        self.module_patcher = patch.dict(
            sys.modules,
            {
                "psycopg": psycopg,
                "psycopg.rows": psycopg_rows,
                "app.config": config,
                "app.business_rules": business_rules,
            },
        )
        self.module_patcher.start()

        def cleanup() -> None:
            for name in ("app.excel_export", "app.digest", "app.lead_notes"):
                sys.modules.pop(name, None)
            self.module_patcher.stop()
            for name, module in self.previous_modules.items():
                if module is not None:
                    sys.modules[name] = module

        self.addCleanup(cleanup)
        self.excel_export = importlib.import_module("app.excel_export")

    def assert_plain_excel_package(self, content: bytes) -> None:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            names = archive.namelist()
            self.assertFalse(
                any(name.startswith("xl/tables/") for name in names),
                f"Excel package contains table XML files: {names}",
            )

            xml_names = [
                name
                for name in names
                if name.endswith(".xml") or name.endswith(".rels")
            ]
            for name in xml_names:
                data = archive.read(name)
                self.assertIsNone(
                    INVALID_XML_CONTROL_BYTES_RE.search(data),
                    f"{name} contains invalid XML control characters",
                )
                ET.fromstring(data)

            external_hyperlink_rels = []
            for name in xml_names:
                if not name.endswith(".rels"):
                    continue
                root = ET.fromstring(archive.read(name))
                for relation in root:
                    attributes = relation.attrib
                    if (
                        attributes.get("Type", "").endswith("/hyperlink")
                        or attributes.get("TargetMode") == "External"
                    ):
                        external_hyperlink_rels.append((name, attributes))

        self.assertEqual(external_hyperlink_rels, [])

    def assert_export_workbook_shape(self, content: bytes, sheet_name: str) -> None:
        workbook = load_workbook(io.BytesIO(content))
        self.assertIn(sheet_name, workbook.sheetnames)
        sheet = workbook[sheet_name]
        headers = [cell.value for cell in sheet[1]]
        all_text = "\n".join(
            str(cell.value or "")
            for row_cells in sheet.iter_rows()
            for cell in row_cells
        )

        self.assertNotIn(FORBIDDEN_COPY_COLUMN, headers)
        self.assertNotIn(FORBIDDEN_ROLE_WORD, all_text.lower())
        for column in KEY_LEAD_COLUMNS:
            self.assertIn(column, headers)
        self.assertEqual(list(sheet.tables.values()), [])
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(
            sheet.auto_filter.ref,
            f"A1:{get_column_letter(len(self.excel_export.EXCEL_EXPORT_COLUMNS))}{sheet.max_row}",
        )
        for row_cells in sheet.iter_rows():
            for cell in row_cells:
                self.assertNotEqual(cell.data_type, "f")

        self.assert_plain_excel_package(content)

    def test_today_excel_workbook_has_expected_columns(self) -> None:
        now = datetime(2026, 5, 5, 12, 0, tzinfo=timezone.utc)
        source_rows = [make_row("today-1", deadline_at=now + timedelta(days=4))]

        with patch.object(
            self.excel_export.digest,
            "select_operational_lead_queue_rows",
            return_value=(source_rows, SimpleNamespace()),
        ):
            rows = self.excel_export.build_excel_export_rows(
                self.excel_export.EXCEL_EXPORT_TODAY,
                now=now,
            )

        self.assertNotIn("Карточка для копирования", rows[0])
        content = self.excel_export.build_excel_workbook_bytes(
            sheet_name="Новые клиентские лиды",
            rows=rows,
        )
        workbook = load_workbook(io.BytesIO(content))
        sheet = workbook["Новые клиентские лиды"]

        expected_columns = [
            "№",
            "ID закупки / external_id",
            "Название",
            "Заказчик",
            "НМЦК",
            "Срок подачи",
            "Приоритет",
            "Уверенность",
            "Документы",
            "Рекомендация",
            "Сигнал по заказчику",
            "Гипотеза / история заказчика",
            "Возможные потребности",
            "Кого искать",
            "Контакт закупки",
            "Первая фраза",
            "Что выяснить",
            "Следующее действие",
            "Ссылка",
            "Статус",
            "Примечание менеджера",
        ]
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(self.excel_export.EXCEL_EXPORT_COLUMNS, expected_columns)
        self.assertEqual(headers, expected_columns)
        self.assertNotIn("Карточка для копирования", headers)
        self.assertEqual(sheet["A2"].value, 1)
        self.assertEqual(sheet["B2"].value, "today-1")
        self.assertEqual(sheet["I2"].value, "Документы: ТЗ найдено, высокая уверенность")
        self.assertEqual(sheet["K2"].value, "Заказчик показывает закупочный сигнал по инфраструктуре.")
        self.assertIn("серверы", sheet["M2"].value)
        self.assertIn("ИТ-директор", sheet["N2"].value)
        self.assertEqual(sheet["S2"].value, "https://example.test/tender")
        self.assertIsNone(sheet["S2"].hyperlink)
        self.assertEqual(sheet["T2"].value, "Новый")
        self.assertEqual(rows[0]["Примечание менеджера"], "")
        self.assertEqual(list(sheet.tables.values()), [])
        self.assertNotIn(FORBIDDEN_ROLE_WORD, "\n".join(str(cell.value or "") for cell in sheet[2]).lower())
        self.assert_plain_excel_package(content)

    def test_excel_export_includes_latest_manager_note_safely(self) -> None:
        now = datetime(2026, 5, 5, 12, 0, tzinfo=timezone.utc)
        source_rows = [make_row("today-1", deadline_at=now + timedelta(days=4))]

        with (
            patch.object(
                self.excel_export.digest,
                "select_operational_lead_queue_rows",
                return_value=(source_rows, SimpleNamespace()),
            ),
            patch.object(
                self.excel_export,
                "fetch_latest_lead_notes_for_tenders",
                return_value={"today-1": "=проверить контакт\nпосле звонка"},
            ),
        ):
            rows = self.excel_export.build_excel_export_rows(
                self.excel_export.EXCEL_EXPORT_TODAY,
                now=now,
            )

        self.assertEqual(rows[0]["Примечание менеджера"], "=проверить контакт\nпосле звонка")
        content = self.excel_export.build_excel_workbook_bytes(
            sheet_name="Новые клиентские лиды",
            rows=rows,
        )
        workbook = load_workbook(io.BytesIO(content))
        sheet = workbook["Новые клиентские лиды"]

        self.assertEqual(sheet["U2"].value, "'=проверить контакт\nпосле звонка")
        self.assertNotEqual(sheet["U2"].data_type, "f")
        self.assert_plain_excel_package(content)

    def test_unprocessed_excel_export_uses_only_visible_operational_go_leads(self) -> None:
        now = datetime(2026, 5, 5, 12, 0, tzinfo=timezone.utc)
        source_rows = [
            make_row("all-1", deadline_at=now + timedelta(days=4)),
            make_row(
                "all-in-work",
                deadline_at=now + timedelta(days=5),
                latest_sales_status="sales_in_work",
                latest_sales_status_created_at=now - timedelta(minutes=5),
            ),
            make_row("all-not-relevant", deadline_at=now + timedelta(days=5), latest_sales_status="sales_not_relevant"),
            make_row("all-skip", deadline_at=now + timedelta(days=5), latest_sales_status="skip"),
            make_row("all-returned", deadline_at=now + timedelta(days=5), latest_sales_status="sales_returned"),
        ]

        with patch.object(
            self.excel_export.digest,
            "get_operational_lead_report_rows",
            return_value=source_rows,
        ) as selector:
            rows = self.excel_export.build_excel_export_rows(
                self.excel_export.EXCEL_EXPORT_ALL,
                now=now,
            )

        selector.assert_called_once_with(limit=None)
        self.assertEqual({row["ID закупки / external_id"] for row in rows}, {"all-1", "all-returned"})
        statuses = {row["ID закупки / external_id"]: row["Статус"] for row in rows}
        self.assertEqual(statuses["all-1"], "Новый")
        self.assertEqual(statuses["all-returned"], "Новый")
        self.assertNotIn("В работе", set(statuses.values()))

        content = self.excel_export.build_excel_workbook_bytes(
            sheet_name="Необработанные клиентские лиды",
            rows=rows,
            table_name="LeadAllTable",
        )
        workbook = load_workbook(io.BytesIO(content))
        sheet = workbook["Необработанные клиентские лиды"]
        headers = [cell.value for cell in sheet[1]]
        self.assertNotIn("Карточка для копирования", headers)
        self.assertEqual(sheet["T2"].value, "Новый")
        self.assertEqual(sheet["T3"].value, "Новый")
        self.assertNotIn(
            FORBIDDEN_ROLE_WORD,
            "\n".join(str(cell.value or "") for row in sheet.iter_rows() for cell in row).lower(),
        )
        self.assertEqual(list(sheet.tables.values()), [])
        self.assert_plain_excel_package(content)

    def test_in_work_excel_export_includes_latest_sales_in_work(self) -> None:
        now = datetime(2026, 5, 5, 12, 0, tzinfo=timezone.utc)
        source_rows = [
            make_row(
                "in-work",
                deadline_at=now + timedelta(days=4),
                latest_sales_status="sales_in_work",
                latest_sales_status_created_at=now - timedelta(minutes=5),
            )
        ]

        with patch.object(
            self.excel_export.digest,
            "select_operational_lead_in_work_rows",
            return_value=(source_rows, SimpleNamespace()),
        ):
            rows = self.excel_export.build_excel_export_rows(
                self.excel_export.EXCEL_EXPORT_IN_WORK,
                now=now,
            )

        self.assertNotIn("Карточка для копирования", rows[0])
        self.assertEqual(rows[0]["Статус"], "В работе")

        content = self.excel_export.build_excel_workbook_bytes(
            sheet_name="Лиды в работе",
            rows=rows,
            table_name="LeadInWorkTable",
        )
        workbook = load_workbook(io.BytesIO(content))
        sheet = workbook["Лиды в работе"]
        headers = [cell.value for cell in sheet[1]]

        self.assertNotIn("Карточка для копирования", headers)
        for column in (
            "ID закупки / external_id",
            "Гипотеза / история заказчика",
            "Кого искать",
            "Следующее действие",
            "Статус",
        ):
            self.assertIn(column, headers)
        self.assertEqual(sheet["B2"].value, "in-work")
        self.assertEqual(sheet["T2"].value, "В работе")
        self.assertEqual(list(sheet.tables.values()), [])
        self.assert_plain_excel_package(content)

    def test_bitrix_excel_export_uses_in_work_rows_and_manager_note_in_comments(self) -> None:
        now = datetime(2026, 5, 5, 12, 0, tzinfo=timezone.utc)
        source_rows = [
            make_row(
                "bitrix-work",
                deadline_at=now + timedelta(days=4),
                latest_sales_status="sales_in_work",
                latest_sales_status_created_at=now - timedelta(minutes=5),
                raw_full={
                    "orderNumber": "32615959798",
                    "lawName": "223-ФЗ",
                    "contactPerson": "Иван Иванов",
                    "contactPhone": "+7 999 000-00-00",
                    "contactEMail": "buyer@example.test",
                    "purchaseMethodName": "Запрос предложений",
                    "statusName": "Подача заявок",
                    "platformName": "ЕИС",
                    "deliveryPlace": "Москва",
                    "organizer": {
                        "id": "org-42",
                        "fullName": "ПАО Организатор",
                        "inn": "7700000000",
                        "kpp": "770001001",
                        "address": "Москва, ул. Тестовая, 1",
                    },
                },
            )
        ]

        with (
            patch.object(
                self.excel_export.digest,
                "select_operational_lead_queue_rows",
                side_effect=AssertionError("Bitrix export must not use raw queue rows"),
            ),
            patch.object(
                self.excel_export.digest,
                "select_operational_lead_all_rows",
                side_effect=AssertionError("Bitrix export must not use unprocessed rows"),
            ),
            patch.object(
                self.excel_export.digest,
                "select_operational_lead_in_work_rows",
                return_value=(source_rows, SimpleNamespace()),
            ) as in_work_selector,
            patch.object(
                self.excel_export,
                "fetch_latest_lead_notes_for_tenders",
                return_value={"bitrix-work": "=перезвонить после обеда"},
            ),
        ):
            export = self.excel_export.build_excel_export_file(
                self.excel_export.EXCEL_EXPORT_BITRIX_IN_WORK,
                now=now,
            )

        in_work_selector.assert_called_once()
        self.assertEqual(export.rows_count, 1)
        self.assertIn("bitrix_in_work", export.filename)
        workbook = load_workbook(io.BytesIO(export.content))
        sheet = workbook["Bitrix лиды"]
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(headers, self.excel_export.BITRIX_EXPORT_COLUMNS)
        self.assertEqual(sheet["A2"].value.startswith("ООО Заказчик / Закупка bitrix-work"), True)
        self.assertEqual(sheet["C2"].value, "Иван Иванов")
        self.assertEqual(sheet["D2"].value, "+7 999 000-00-00")
        self.assertEqual(sheet["E2"].value, "buyer@example.test")
        self.assertEqual(sheet["H2"].value, "NEW")
        self.assertEqual(sheet["I2"].value, "1")
        self.assertEqual(sheet["K2"].value, 1250000)
        self.assertEqual(sheet["M2"].value, "Y")
        self.assertEqual(sheet["N2"].value, "bitrix-work")
        self.assertEqual(sheet["Q2"].value, "2026-05-09 12:00:00")
        self.assertEqual(sheet["T2"].value, "32615959798")
        self.assertEqual(sheet["U2"].value, "143")
        self.assertEqual(sheet["V2"].value, "ЕИС")
        self.assertIn("Примечание менеджера:\n=перезвонить после обеда", sheet["G2"].value)
        self.assertIn("Следующее действие:", sheet["G2"].value)
        self.assertNotEqual(sheet["G2"].data_type, "f")
        self.assertEqual(sheet.auto_filter.ref, f"A1:{get_column_letter(len(headers))}2")
        self.assert_plain_excel_package(export.content)

    def test_today_and_unprocessed_excel_hide_in_work_and_negative_feedback(self) -> None:
        now = datetime(2026, 5, 5, 12, 0, tzinfo=timezone.utc)
        source_rows = [
            make_row("visible", deadline_at=now + timedelta(days=4)),
            make_row("expired", deadline_at=now - timedelta(hours=1)),
            make_row("too-soon", deadline_at=now + timedelta(days=2), llm_report_recommendation="maybe"),
            make_row("in-work", deadline_at=now + timedelta(days=4), latest_sales_status="sales_in_work"),
            make_row("not-relevant", deadline_at=now + timedelta(days=4), latest_sales_status="sales_not_relevant"),
            make_row("returned", deadline_at=now + timedelta(days=4), latest_sales_status="sales_returned"),
        ]

        with patch.object(
            self.excel_export.digest,
            "get_operational_lead_report_rows",
            return_value=source_rows,
        ):
            today_rows = self.excel_export.build_excel_export_rows(
                self.excel_export.EXCEL_EXPORT_TODAY,
                now=now,
            )

        with patch.object(
            self.excel_export.digest,
            "get_operational_lead_report_rows",
            return_value=source_rows,
        ):
            all_rows = self.excel_export.build_excel_export_rows(
                self.excel_export.EXCEL_EXPORT_ALL,
                now=now,
            )

        self.assertEqual(
            {row["ID закупки / external_id"] for row in today_rows},
            {"visible", "expired", "too-soon", "returned"},
        )
        self.assertEqual(
            {row["ID закупки / external_id"] for row in all_rows},
            {"visible", "expired", "too-soon", "returned"},
        )
        self.assertNotIn("В работе", {row["Статус"] for row in all_rows})
        returned_row = next(row for row in all_rows if row["ID закупки / external_id"] == "returned")
        self.assertEqual(returned_row["Статус"], "Новый")

    def test_today_excel_export_uses_current_lead_queue(self) -> None:
        now = datetime(2026, 5, 5, 12, 0, tzinfo=timezone.utc)
        source_rows = [
            make_row("lead-queue", deadline_at=now + timedelta(days=4))
        ]

        with patch.object(
            self.excel_export.digest,
            "select_operational_lead_queue_rows",
            return_value=(source_rows, SimpleNamespace()),
        ) as selector:
            rows = self.excel_export.build_excel_export_rows(
                self.excel_export.EXCEL_EXPORT_TODAY,
                now=now,
            )

        selector.assert_called_once()
        self.assertEqual(rows[0]["ID закупки / external_id"], "lead-queue")

    def test_excel_today_all_and_in_work_match_operational_selections(self) -> None:
        now = datetime(2026, 5, 5, 12, 0, tzinfo=timezone.utc)
        today_created = datetime(2026, 5, 5, 8, 0, tzinfo=timezone.utc)
        yesterday_moscow = datetime(2026, 5, 4, 20, 59, tzinfo=timezone.utc)
        source_rows = [
            make_row(
                "today-triage-go-report-maybe",
                deadline_at=now + timedelta(days=4),
                llm_report_recommendation="maybe",
                lead_triage_decision="go",
                lead_triage_created_at=today_created - timedelta(minutes=5),
                llm_report_created_at=today_created,
            ),
            make_row(
                "yesterday-triage-go-report-maybe",
                deadline_at=now + timedelta(days=4),
                llm_report_recommendation="maybe",
                lead_triage_decision="go",
                lead_triage_created_at=yesterday_moscow - timedelta(minutes=5),
                llm_report_created_at=yesterday_moscow,
            ),
            make_row(
                "today-triage-maybe",
                deadline_at=now + timedelta(days=4),
                llm_report_recommendation="maybe",
                lead_triage_decision="maybe",
                llm_report_created_at=today_created,
            ),
            make_row(
                "today-sales-in-work",
                deadline_at=now + timedelta(days=4),
                llm_report_recommendation="maybe",
                lead_triage_decision="go",
                lead_triage_created_at=today_created - timedelta(minutes=5),
                latest_sales_status="sales_in_work",
                llm_report_created_at=today_created,
            ),
            make_row(
                "legacy-report-go-today",
                deadline_at=now + timedelta(days=4),
                llm_report_recommendation="go",
                llm_report_created_at=today_created,
            ),
            make_row(
                "legacy-report-go-yesterday",
                deadline_at=now + timedelta(days=4),
                llm_report_recommendation="go",
                llm_report_created_at=yesterday_moscow,
            ),
        ]

        with patch.object(
            self.excel_export.digest,
            "get_operational_lead_report_rows",
            return_value=source_rows,
        ):
            today_rows = self.excel_export.build_excel_export_rows(
                self.excel_export.EXCEL_EXPORT_TODAY,
                now=now,
            )

        with patch.object(
            self.excel_export.digest,
            "get_operational_lead_report_rows",
            return_value=source_rows,
        ):
            all_rows = self.excel_export.build_excel_export_rows(
                self.excel_export.EXCEL_EXPORT_ALL,
                now=now,
            )

        with patch.object(
            self.excel_export.digest,
            "get_operational_lead_report_in_work_rows",
            return_value=source_rows,
        ):
            in_work_rows = self.excel_export.build_excel_export_rows(
                self.excel_export.EXCEL_EXPORT_IN_WORK,
                now=now,
            )

        id_column = "ID закупки / external_id"
        self.assertEqual(
            {row[id_column] for row in today_rows},
            {
                "today-triage-go-report-maybe",
                "today-triage-maybe",
                "legacy-report-go-today",
            },
        )
        self.assertEqual(
            {row[id_column] for row in all_rows},
            {
                "today-triage-go-report-maybe",
                "yesterday-triage-go-report-maybe",
                "today-triage-maybe",
                "legacy-report-go-today",
                "legacy-report-go-yesterday",
            },
        )
        self.assertEqual(
            {row[id_column] for row in in_work_rows},
            {"today-sales-in-work"},
        )

    def test_excel_exports_use_operational_queue_decision(self) -> None:
        now = datetime(2026, 5, 5, 12, 0, tzinfo=timezone.utc)
        source_rows = [
            make_row(
                "triage-go-report-maybe",
                deadline_at=now + timedelta(days=4),
                llm_report_recommendation="maybe",
                lead_triage_decision="go",
            ),
            make_row(
                "triage-maybe-report-maybe",
                deadline_at=now + timedelta(days=4),
                llm_report_recommendation="maybe",
                lead_triage_decision="maybe",
            ),
            make_row(
                "triage-reject-report-go",
                deadline_at=now + timedelta(days=4),
                llm_report_recommendation="go",
                lead_triage_decision="reject",
                lead_triage_created_at=now + timedelta(minutes=1),
            ),
            make_row("legacy-report-go", deadline_at=now + timedelta(days=4)),
            make_row(
                "triage-go-in-work",
                deadline_at=now + timedelta(days=4),
                llm_report_recommendation="maybe",
                lead_triage_decision="go",
                latest_sales_status="sales_in_work",
            ),
            make_row(
                "triage-go-negative",
                deadline_at=now + timedelta(days=4),
                llm_report_recommendation="maybe",
                lead_triage_decision="go",
                latest_sales_status="sales_not_relevant",
            ),
        ]

        with patch.object(
            self.excel_export.digest,
            "get_operational_lead_report_rows",
            return_value=source_rows,
        ):
            today_rows = self.excel_export.build_excel_export_rows(
                self.excel_export.EXCEL_EXPORT_TODAY,
                now=now,
            )

        with patch.object(
            self.excel_export.digest,
            "get_operational_lead_report_rows",
            return_value=source_rows,
        ):
            all_rows = self.excel_export.build_excel_export_rows(
                self.excel_export.EXCEL_EXPORT_ALL,
                now=now,
            )

        id_column = self.excel_export.EXCEL_EXPORT_COLUMNS[1]
        recommendation_column = "Рекомендация"
        expected_ids = {"triage-go-report-maybe", "triage-maybe-report-maybe", "legacy-report-go"}
        self.assertEqual({row[id_column] for row in today_rows}, expected_ids)
        self.assertEqual({row[id_column] for row in all_rows}, expected_ids)

        recommendations = {row[id_column]: row[recommendation_column] for row in all_rows}
        self.assertEqual(
            recommendations["triage-go-report-maybe"],
            self.excel_export.digest.call_first_recommendation_text("go"),
        )
        self.assertNotEqual(
            recommendations["triage-go-report-maybe"],
            self.excel_export.digest.call_first_recommendation_text("maybe"),
        )

        content = self.excel_export.build_excel_workbook_bytes(
            sheet_name="Operational Leads",
            rows=all_rows,
            table_name="OperationalLeadsTable",
        )
        workbook = load_workbook(io.BytesIO(content))
        headers = [cell.value for cell in workbook["Operational Leads"][1]]
        self.assertNotIn(FORBIDDEN_COPY_COLUMN, headers)
        self.assertNotIn(FORBIDDEN_COPY_COLUMN, all_rows[0])

    def test_excel_export_hides_latest_lead_triage_reject_without_full_report(self) -> None:
        now = datetime(2026, 5, 5, 12, 0, tzinfo=timezone.utc)
        rejected = make_row(
            "triage-reject-without-full-report",
            deadline_at=now + timedelta(days=4),
            lead_triage_decision="reject",
            lead_triage_created_at=now,
        )
        rejected["llm_report_analysis_type"] = None
        rejected["llm_report_created_at"] = None
        rejected["llm_report_result"] = {}
        visible = make_row("visible", deadline_at=now + timedelta(days=4))

        with patch.object(
            self.excel_export.digest,
            "select_operational_lead_all_rows",
            return_value=([rejected, visible], SimpleNamespace()),
        ):
            rows = self.excel_export.build_excel_export_rows(
                self.excel_export.EXCEL_EXPORT_ALL,
                now=now,
            )

        id_column = self.excel_export.EXCEL_EXPORT_COLUMNS[1]
        self.assertEqual([row[id_column] for row in rows], ["visible"])

    def test_excel_export_hides_current_hard_noise_from_today_and_all(self) -> None:
        now = datetime(2026, 5, 5, 12, 0, tzinfo=timezone.utc)
        noise = make_row("hard-noise", deadline_at=now + timedelta(days=4))
        noise["title"] = "Kaspersky antivirus license renewal"
        visible = make_row("visible", deadline_at=now + timedelta(days=4))
        source_rows = [noise, visible]

        with patch.object(
            self.excel_export.digest,
            "get_operational_lead_report_rows",
            return_value=source_rows,
        ):
            today_rows = self.excel_export.build_excel_export_rows(
                self.excel_export.EXCEL_EXPORT_TODAY,
                now=now,
            )

        with patch.object(
            self.excel_export.digest,
            "get_operational_lead_report_rows",
            return_value=source_rows,
        ):
            all_rows = self.excel_export.build_excel_export_rows(
                self.excel_export.EXCEL_EXPORT_ALL,
                now=now,
            )

        id_column = self.excel_export.EXCEL_EXPORT_COLUMNS[1]
        self.assertEqual([row[id_column] for row in today_rows], ["visible"])
        self.assertEqual([row[id_column] for row in all_rows], ["visible"])

    def test_excel_all_matches_telegram_all_leads_filtering(self) -> None:
        now = datetime(2026, 5, 5, 12, 0, tzinfo=timezone.utc)
        visible = make_row("visible", deadline_at=now + timedelta(days=4))
        hard_noise = make_row("hard-noise", deadline_at=now + timedelta(days=4))
        hard_noise["title"] = "Kaspersky antivirus license renewal"
        in_work = make_row(
            "in-work",
            deadline_at=now + timedelta(days=4),
            latest_sales_status=self.excel_export.digest.SALES_IN_WORK_STATUS,
        )
        negative = make_row(
            "negative",
            deadline_at=now + timedelta(days=4),
            latest_sales_status=self.excel_export.digest.SALES_NOT_RELEVANT_STATUS,
        )
        source_rows = [hard_noise, in_work, negative, visible]

        with patch.object(
            self.excel_export.digest,
            "get_operational_lead_report_rows",
            return_value=source_rows,
        ):
            telegram_cards, telegram_stats = (
                self.excel_export.digest.build_operational_lead_all_cards_with_stats(
                    limit=None,
                    now=now,
                )
            )

        with patch.object(
            self.excel_export.digest,
            "get_operational_lead_report_rows",
            return_value=source_rows,
        ):
            excel_rows = self.excel_export.build_excel_export_rows(
                self.excel_export.EXCEL_EXPORT_ALL,
                now=now,
            )

        id_column = self.excel_export.EXCEL_EXPORT_COLUMNS[1]
        self.assertEqual([card.external_id for card in telegram_cards], ["visible"])
        self.assertEqual([row[id_column] for row in excel_rows], ["visible"])
        self.assertEqual(telegram_stats.hidden_by_lead_hard_noise, 1)
        self.assertEqual(telegram_stats.hidden_in_work, 1)
        self.assertEqual(telegram_stats.hidden_negative_feedback, 1)

    def test_excel_all_dedupes_duplicate_operational_lead_reports(self) -> None:
        now = datetime(2026, 5, 5, 12, 0, tzinfo=timezone.utc)
        older = make_row(
            "older-report-row",
            deadline_at=now + timedelta(days=4),
            llm_report_created_at=now - timedelta(hours=1),
        )
        newer = make_row(
            "newer-report-row",
            deadline_at=now + timedelta(days=4),
            llm_report_created_at=now,
        )
        older["external_id"] = "same-external"
        newer["external_id"] = "same-external"

        with patch.object(
            self.excel_export.digest,
            "get_operational_lead_report_rows",
            return_value=[older, newer],
        ):
            rows = self.excel_export.build_excel_export_rows(
                self.excel_export.EXCEL_EXPORT_ALL,
                now=now,
            )

        id_column = self.excel_export.EXCEL_EXPORT_COLUMNS[1]
        title_column = self.excel_export.EXCEL_EXPORT_COLUMNS[2]
        self.assertEqual([row[id_column] for row in rows], ["same-external"])
        self.assertIn("newer-report-row", rows[0][title_column])

    def test_phone_email_and_url_are_separate_columns(self) -> None:
        now = datetime(2026, 5, 5, 12, 0, tzinfo=timezone.utc)
        row = make_row(
            "contacts",
            deadline_at=now + timedelta(days=4),
            raw_full={
                "orderNumber": "REG-42",
                "contactPerson": "Иван Иванов",
                "contactPhone": "+7 999 000-00-00",
                "contactEMail": "buyer@example.test",
                "tenderUrl": "https://example.test/reg-42",
            },
        )

        export_row = self.excel_export.excel_row_from_tender(row, now=now)

        self.assertEqual(export_row["ID закупки / external_id"], "contacts")
        self.assertEqual(export_row["Контакт закупки"], "Иван Иванов, +7 999 000-00-00, buyer@example.test")
        self.assertEqual(export_row["Ссылка"], "https://example.test/tender")

    def test_export_files_open_as_plain_worksheets_without_table_xml(self) -> None:
        now = datetime(2026, 5, 5, 12, 0, tzinfo=timezone.utc)
        rows_by_mode = {
            self.excel_export.EXCEL_EXPORT_TODAY: [make_row("today", deadline_at=now + timedelta(days=4))],
            self.excel_export.EXCEL_EXPORT_ALL: [
                make_row("all-new", deadline_at=now + timedelta(days=4)),
            ],
            self.excel_export.EXCEL_EXPORT_IN_WORK: [
                make_row("work", deadline_at=now + timedelta(days=4), latest_sales_status="sales_in_work")
            ],
        }

        def queue_selector(*args, **kwargs):
            return rows_by_mode[self.excel_export.EXCEL_EXPORT_TODAY], SimpleNamespace()

        def all_selector(*args, **kwargs):
            return rows_by_mode[self.excel_export.EXCEL_EXPORT_ALL], SimpleNamespace()

        def in_work_selector(*args, **kwargs):
            return rows_by_mode[self.excel_export.EXCEL_EXPORT_IN_WORK], SimpleNamespace()

        with (
            patch.object(self.excel_export.digest, "select_operational_lead_queue_rows", side_effect=queue_selector),
            patch.object(self.excel_export.digest, "select_operational_lead_all_rows", side_effect=all_selector),
            patch.object(self.excel_export.digest, "select_operational_lead_in_work_rows", side_effect=in_work_selector),
        ):
            exports = {
                mode: self.excel_export.build_excel_export_file(mode, now=now)
                for mode in rows_by_mode
            }

        expected_sheets = {
            self.excel_export.EXCEL_EXPORT_TODAY: "Новые клиентские лиды",
            self.excel_export.EXCEL_EXPORT_ALL: "Необработанные клиентские лиды",
            self.excel_export.EXCEL_EXPORT_IN_WORK: "Лиды в работе",
        }

        for mode, export in exports.items():
            self.assert_export_workbook_shape(export.content, expected_sheets[mode])

    def test_safe_cell_values_keep_urls_plain_and_remove_excel_repair_risks(self) -> None:
        row = {column: "" for column in self.excel_export.EXCEL_EXPORT_COLUMNS}
        row.update(
            {
                "№": 1,
                "ID закупки / external_id": "=external-id",
                "Название": "line one\r\nline two\rline three\x00\x1F",
                "Возможные потребности": ["=first item", "second\x01 item"],
                "Следующее действие": "x" * 33050,
                "Ссылка": "https://example.test/tender?x=1",
                "Статус": "@new",
            }
        )
        content = self.excel_export.build_excel_workbook_bytes(
            sheet_name="Новые клиентские лиды",
            rows=[row],
        )

        workbook = load_workbook(io.BytesIO(content))
        sheet = workbook["Новые клиентские лиды"]
        self.assertEqual(sheet["B2"].value, "'=external-id")
        self.assertEqual(sheet["C2"].value, "line one\nline two\nline three")
        self.assertEqual(sheet["M2"].value, "'=first item\nsecond item")
        self.assertEqual(len(sheet["R2"].value), 32000)
        self.assertEqual(sheet["S2"].value, "https://example.test/tender?x=1")
        self.assertIsNone(sheet["S2"].hyperlink)
        self.assertEqual(sheet["T2"].value, "'@new")
        for row_cells in sheet.iter_rows():
            for cell in row_cells:
                self.assertNotEqual(cell.data_type, "f")
        self.assert_plain_excel_package(content)

    def test_empty_workbook_has_header_but_no_broken_table(self) -> None:
        content = self.excel_export.build_excel_workbook_bytes(
            sheet_name="Новые клиентские лиды",
            rows=[],
            table_name="LeadQueueTable",
        )

        workbook = load_workbook(io.BytesIO(content))
        sheet = workbook["Новые клиентские лиды"]
        self.assertEqual(sheet.max_row, 1)
        self.assertEqual([cell.value for cell in sheet[1]], self.excel_export.EXCEL_EXPORT_COLUMNS)
        self.assertEqual(list(sheet.tables.values()), [])
        self.assert_plain_excel_package(content)


if __name__ == "__main__":
    unittest.main()
