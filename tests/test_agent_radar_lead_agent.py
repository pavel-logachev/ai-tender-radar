"""All customer data is synthetic; phone area code 000 and INN 0000000000 are invalid.

Offline tests for the prompt-driven lead research agent: tools, loop, grounding, card, store, delivery."""
from __future__ import annotations

import asyncio
import json
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import httpx

from agent_radar.lead_agent import agent as A
from agent_radar.lead_agent import card, delivery, export, feedback, grounding, health, runner, tools
from agent_radar.lead_agent.store import LeadStore

NOW = datetime(2026, 10, 1, 9, 0, tzinfo=timezone.utc)


def make_row(tender_id="example-source:1", status="proposal", deadline="2026-10-05T09:00:00Z"):
    return {"card": {"id": tender_id, "title": "Поставка серверов", "customer_name": "ООО Учебный заказчик",
                     "source_url": "https://tenders.example.org/process/light/x", "description": "Нужны 2 сервера",
                     "documents": [{"id": "d1", "title": "ТЗ", "text": "Контакт: Иванов Иван тел. +7 (000) 000-00-00"}]},
            "opportunity": {"buyer": {"inn": "0000000000"}, "status": status, "publication_date": "2026-09-20T00:00:00Z",
                            "acceptance_end_date": deadline}}


GOOD = {"verdict": "lead", "grade": "A", "one_line": "Заказчик покупает 2 сервера",
        "signal": {"what_they_buy": "2 сервера", "deadline": "05.10.2026"},
        "customer": {"name": "ООО Учебный заказчик", "economics_short": "Крупный, выручка 1 млрд"},
        "contacts": [{"name": "Иванов Иван", "role": "закупки", "phone": "+7 (000) 000-00-00", "email": None}],
        "talk_track": {"opening": "Здравствуйте", "hooks": ["a", "b"], "questions": ["q"]},
        "important": None, "gaps": ["ИТ-руководитель не найден"]}
READ = "Контакт: Иванов Иван тел. +7 (000) 000-00-00"


class ToolsTest(unittest.TestCase):
    def test_non_public_and_unsafe_urls_are_refused(self):
        for url in ("http://127.0.0.1/", "http://10.0.0.5/x", "http://[::1]/", "http://169.254.169.254/latest",
                    "file:///etc/passwd", "ftp://example.com/", "https://user:" + "pw@example.com/"):
            with self.subTest(url=url):
                self.assertIsNotNone(tools.public_problem(url))
        self.assertTrue(tools.fetch_page("http://127.0.0.1/").startswith("refused"))

    def test_every_redirect_hop_is_checked(self):
        def handler(request):
            return httpx.Response(302, headers={"location": "http://10.0.0.1/secret"})
        checks = iter([None, "non-public address"])
        with patch.object(tools, "public_problem", side_effect=lambda url: next(checks)):
            out = tools.fetch_page("http://example.com/", transport=httpx.MockTransport(handler))
        self.assertEqual(out, "refused: non-public address")

    def test_fetch_returns_marked_text_and_links(self):
        html = "<html><head><title>t</title></head><body><p>Закупки: <a href='/c'>контакты</a></p><script>x()</script></body></html>"
        handler = lambda request: httpx.Response(200, headers={"content-type": "text/html; charset=utf-8"}, content=html.encode())
        with patch.object(tools, "public_problem", return_value=None):
            out = tools.fetch_page("https://example.com/p", with_links=True, transport=httpx.MockTransport(handler))
        self.assertIn(tools.UNTRUSTED.strip(), out)
        self.assertIn("Закупки:", out)
        self.assertNotIn("x()", out)
        self.assertIn("https://example.com/c", out)

    def test_fns_output_keeps_latest_years_compact(self):
        reports = [{"period": str(year), "published": True, "isCb": False, "gainSum": year * 10,
                    "typeCorrections": [{"correction": {"financialResult": {"current2110": year * 10}, "balance": {}}}]}
                   for year in (2023, 2025, 2024)]
        data = tools.compact_fns({"id": 5, "inn": "x", "shortName": "АО"}, reports)
        self.assertEqual([row["period"] for row in data["years"]], ["2023", "2024", "2025"])
        self.assertEqual(data["years"][-1]["revenue_2110"], 20250)
        self.assertLess(len(json.dumps(data)), 3000)

    def test_procedure_lists_its_documents_and_reads_them(self):
        procedure = tools.TenderTools(make_row())
        self.assertEqual([doc["doc_id"] for doc in json.loads(procedure.read_procedure())["documents"]], ["d1"])
        self.assertEqual(procedure.read_document("d2"), "no such document")
        self.assertIn("Иванов", procedure.read_document("d1"))


def scripted(*messages, cost=0.002):
    sent = []
    queue = list(messages)

    def handler(request):
        sent.append(json.loads(request.content))
        message = queue.pop(0) if len(queue) > 1 else queue[0]
        return httpx.Response(200, json={"choices": [{"message": message}], "usage": {"cost": cost, "prompt_tokens": 10, "completion_tokens": 5}})
    return httpx.MockTransport(handler), sent


def tool_call(name, args, call_id="c1"):
    return {"role": "assistant", "content": None, "tool_calls": [{"id": call_id, "type": "function",
            "function": {"name": name, "arguments": json.dumps(args)}}]}


class AgentLoopTest(unittest.TestCase):
    def test_reads_then_finishes_and_accumulates_cost_and_reads(self):
        transport, sent = scripted(tool_call("read_document", {"doc_id": "d1"}),
                                   {"role": "assistant", "content": "```json\n" + json.dumps(GOOD) + "\n```"})
        run = A.run_agent(make_row(), api_key="k", system="sys", transport=transport, sleep=lambda s: None)
        self.assertEqual(A.extract_json(run.final_text)["grade"], "A")
        self.assertEqual((run.tool_calls, run.cost_usd, run.tokens_in), (1, 0.004, 20))
        self.assertIn("Иванов Иван", run.tool_text)
        self.assertEqual(sent[0]["messages"][0]["content"], "sys")

    def test_budget_nudge_forces_final_answer_without_tools(self):
        transport, sent = scripted(tool_call("read_procedure", {}), tool_call("read_procedure", {}, "c2"),
                                   {"role": "assistant", "content": json.dumps(GOOD)})
        run = A.run_agent(make_row(), api_key="k", system="s", max_turns=3, transport=transport, sleep=lambda s: None)
        self.assertEqual(len(sent), 3)
        self.assertNotIn("tool_choice", sent[0])
        self.assertEqual(sent[-1].get("tool_choice"), "none")
        self.assertTrue(any("Бюджет шагов" in str(m.get("content")) for m in sent[-1]["messages"]))
        self.assertTrue(run.final_text)

    def test_tool_failure_and_unknown_tool_do_not_crash(self):
        transport, _ = scripted(tool_call("nope", {}), {"role": "assistant", "content": json.dumps(GOOD)})
        run = A.run_agent(make_row(), api_key="k", system="s", transport=transport, sleep=lambda s: None,
                          tool_impl={"web_search": lambda **kw: 1 / 0})
        self.assertIsNone(run.error)
        self.assertEqual(run.tool_calls, 1)

    def test_transport_errors_are_retried_then_reported(self):
        def boom(request):
            raise httpx.ConnectError("down")
        run = A.run_agent(make_row(), api_key="k", system="s", transport=httpx.MockTransport(boom), sleep=lambda s: None)
        self.assertIsNone(run.final_text)
        self.assertIn("down", run.error)


class GroundingTest(unittest.TestCase):
    def test_verified_person_with_phone_stays_grade_a(self):
        result = grounding.ground_result(json.loads(json.dumps(GOOD)), READ)
        self.assertEqual(result["grade"], "A")
        self.assertTrue(result["contacts"][0]["verified"]["phone"])

    def test_invented_phone_and_email_are_removed_and_grade_recomputed(self):
        bad = json.loads(json.dumps(GOOD))
        bad["contacts"][0].update(phone="+7 (000) 000-11-22", email="ivan@example.org")
        result = grounding.ground_result(bad, READ)
        contact = result["contacts"][0]
        self.assertIsNone(contact["phone"])
        self.assertIsNone(contact["email"])
        self.assertEqual(sorted(contact["removed_unverified"]), ["email", "phone"])
        self.assertEqual((result["grade"], result["verdict"]), ("C", "candidate"))
        self.assertEqual(result["grade_adjusted"], {"from": "A", "to": "C"})

    def test_low_confidence_aggregator_contact_never_earns_grade_a(self):
        data = json.loads(json.dumps(GOOD))
        data["contacts"][0]["confidence"] = "low"
        result = grounding.ground_result(data, READ)
        self.assertEqual(result["grade"], "B")

    def test_name_not_literally_in_sources_is_dropped(self):
        bad = json.loads(json.dumps(GOOD))
        bad["contacts"][0]["name"] = "Примерова Елена"  # derived from an e-mail address, not read anywhere
        result = grounding.ground_result(bad, READ)
        self.assertIsNone(result["contacts"][0]["name"])
        self.assertEqual(result["grade"], "B")  # verified phone without a name

    def test_named_person_with_email_only_is_not_a_ready_lead(self):
        data = json.loads(json.dumps(GOOD))
        data["contacts"] = [{"name": "Иванов Иван", "role": "закупки", "phone": None, "email": "i.ivanov@example.org"}]
        result = grounding.ground_result(data, "Иванов Иван, почта i.ivanov@example.org")
        self.assertEqual((result["grade"], result["verdict"]), ("C", "candidate"))
        self.assertFalse(card.deliverable(result))

    def test_case_endings_and_extensions_are_tolerated(self):
        data = json.loads(json.dumps(GOOD))
        data["contacts"][0].update(name="Иванова Мария", phone="+7 (000) 000-11-33 вн.228; +7 (000) 000-22-44")
        text = "Контакт Ивановой Марии: тел. +7 (000) 000-11-33 вн.228, сот. 8 000 000-22-44"
        result = grounding.ground_result(data, text)
        self.assertEqual(result["grade"], "A")
        self.assertEqual(result["contacts"][0]["phone"].count("+7"), 2)

    def test_candidate_with_a_verified_phone_becomes_a_lead_unless_off_profile(self):
        data = json.loads(json.dumps(GOOD))
        data.update(verdict="candidate", grade="C")
        data["contacts"] = [{"name": None, "role": "Приёмная", "phone": "+7 (000) 000-00-00", "email": None}]
        result = grounding.ground_result(json.loads(json.dumps(data)), READ)
        self.assertEqual((result["verdict"], result["grade"]), ("lead", "B"))
        self.assertEqual(result["verdict_adjusted"], {"from": "candidate", "to": "lead"})
        self.assertTrue(card.deliverable(result))
        data["signal"]["profile_fit"] = False
        off_profile = grounding.ground_result(data, READ)
        self.assertEqual(off_profile["verdict"], "candidate")
        self.assertFalse(card.deliverable(off_profile))


class CardTest(unittest.TestCase):
    def test_card_contains_contact_and_escapes_html(self):
        data = json.loads(json.dumps(GOOD))
        data["one_line"] = "<script>alert(1)</script> & сервер"
        text = card.render_card(data, "https://tenders.example.org/p")
        self.assertIn("☎ +7 (000) 000-00-00", text)
        self.assertNotIn("<script>", text)
        self.assertIn("&amp;", text)
        self.assertIn('<a href="https://tenders.example.org/p">', text)
        self.assertNotIn("http://", card.render_card(data, "http://insecure"))

    def test_directory_number_is_marked_and_a_contact_with_a_phone_is_never_cut_off(self):
        data = json.loads(json.dumps(GOOD))
        data["contacts"] = [{"name": None, "role": "Почта тендеров", "phone": None, "email": "tender@example.org"},
                            {"name": None, "role": "Справочная", "phone": "+7 (000) 000-00-01", "confidence": "low"},
                            {"name": None, "role": "Отдел B2B", "phone": None, "email": "b2b@example.org"},
                            {"name": None, "role": "Приёмная", "phone": "+7 (000) 000-00-02", "confidence": "high"}]
        text = card.render_card(data)
        self.assertLess(text.index("000-00-02"), text.index("000-00-01"))  # official number first
        self.assertLess(text.index("000-00-01"), text.index("tender@example.org"))
        self.assertNotIn("b2b@example.org", text)  # three contacts at most
        self.assertEqual(text.count("номер из стороннего справочника"), 1)
        self.assertEqual(text.count("<i>"), text.count("</i>"))

    def test_long_card_drops_sections_but_never_breaks_tags(self):
        data = json.loads(json.dumps(GOOD))
        data["talk_track"]["hooks"] = ["х" * 400] * 4
        data["talk_track"]["questions"] = ["ы" * 400] * 3
        data["signal"]["what_they_buy"] = "я" * 2000
        data["customer"]["economics_short"] = "э" * 2000
        text = card.render_card(data, "https://tenders.example.org/p")
        self.assertLessEqual(len(text), card.LIMIT)
        self.assertEqual(text.count("<b>"), text.count("</b>"))

    def test_model_markers_and_non_leads_are_rejected(self):
        data = json.loads(json.dumps(GOOD))
        data["one_line"] = "<think>secret</think>"
        with self.assertRaises(ValueError):
            card.render_card(data)
        self.assertFalse(card.deliverable({"verdict": "candidate", "grade": "B"}))
        self.assertFalse(card.deliverable({"verdict": "lead", "grade": "C"}))
        self.assertTrue(card.deliverable({"verdict": "lead", "grade": "B"}))


class StoreTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.store = LeadStore(Path(self.tmp.name) / "private" / "leads.sqlite3")

    def save(self, tender="t1", fingerprint="f1", ok=True):
        return self.store.save(tender_id=tender, fingerprint=fingerprint, model="m", cost_usd=0.01, tokens_in=1, tokens_out=1,
                               tool_calls=2, result=GOOD, card_html="<b>card</b>" if ok else None, deliverable=ok)

    def test_one_result_per_procurement_version(self):
        self.assertIsNotNone(self.save())
        self.assertIsNone(self.save())
        self.assertTrue(self.store.has("t1", "f1"))
        self.assertIsNotNone(self.save(fingerprint="f2"))  # a changed procurement is investigated again

    def test_claim_is_single_and_uncertain_is_never_redelivered(self):
        lead = self.save()
        self.assertEqual([row["id"] for row in self.store.undelivered(7)], [lead])
        self.assertTrue(self.store.claim(lead, 7))
        self.assertFalse(self.store.claim(lead, 7))
        self.store.record_uncertain(lead, 7)
        self.assertEqual(self.store.undelivered(7), [])
        self.assertEqual([row["id"] for row in self.store.undelivered(8)], [lead])  # another chat is independent

    def test_release_allows_retry_only_for_unsent_claims(self):
        lead = self.save()
        self.store.claim(lead, 7)
        self.store.release(lead, 7)
        self.assertEqual(len(self.store.undelivered(7)), 1)
        self.store.claim(lead, 7)
        self.store.record_sent(lead, 7, 99)
        self.store.release(lead, 7)
        self.assertEqual(self.store.undelivered(7), [])

    def test_feedback_requires_delivery_to_that_chat(self):
        lead = self.save()
        self.assertFalse(self.store.add_feedback(lead, 7, 7, "work"))
        self.store.claim(lead, 7)
        self.assertTrue(self.store.add_feedback(lead, 7, 7, "work"))
        self.assertTrue(self.store.add_feedback(lead, 7, 7, "skip"))
        self.assertEqual(self.store.feedback(lead), ["work", "skip"])
        with self.assertRaises(ValueError):
            self.store.add_feedback(lead, 7, 7, "hack")
        self.assertEqual(self.store.stats()["feedback"], {"work": 1, "skip": 1})

    def test_grade_a_is_delivered_before_grade_b(self):
        self.store.save(tender_id="b", fingerprint="1", model="m", cost_usd=0, tokens_in=0, tokens_out=0, tool_calls=0,
                        result={"verdict": "lead", "grade": "B"}, card_html="b", deliverable=True)
        self.store.save(tender_id="a", fingerprint="1", model="m", cost_usd=0, tokens_in=0, tokens_out=0, tool_calls=0,
                        result={"verdict": "lead", "grade": "A"}, card_html="a", deliverable=True)
        self.assertEqual([row["grade"] for row in self.store.undelivered(1)], ["A", "B"])

    def test_a_changed_procurement_is_known_as_an_update_and_stays_away_after_skip(self):
        first = self.save()
        self.assertFalse(self.store.delivered_before("t1"))
        for chat in (7, 8):
            self.store.claim(first, chat)
            self.store.record_sent(first, chat, 99)
        self.assertTrue(self.store.delivered_before("t1"))
        self.assertFalse(self.store.delivered_before("t2"))
        self.store.add_feedback(first, 7, 7, "skip")
        self.store.add_feedback(first, 8, 8, "work")
        second = self.save(fingerprint="f2")
        self.assertEqual(self.store.undelivered(7), [])  # the manager said "Мимо": no second card
        self.assertEqual([row["id"] for row in self.store.undelivered(8)], [second])
        self.store.add_feedback(first, 7, 7, "work")  # the latest answer counts
        self.assertEqual([row["id"] for row in self.store.undelivered(7)], [second])

    def test_a_card_is_not_delivered_after_its_procurement_closed(self):
        closed = self.store.save(tender_id="closed", fingerprint="f", model="m", cost_usd=0, tokens_in=0, tokens_out=0,
                                 tool_calls=0, result=GOOD, card_html="c", deliverable=True, deadline="2026-10-04T09:00:00Z")
        running = self.store.save(tender_id="open", fingerprint="f", model="m", cost_usd=0, tokens_in=0, tokens_out=0,
                                  tool_calls=0, result=GOOD, card_html="o", deliverable=True, deadline="2026-10-06T09:00:00Z")
        undated = self.save()
        monday = "2026-10-05T06:00:00Z"  # first slot after a weekend without deliveries
        self.assertEqual([row["id"] for row in self.store.undelivered(7, open_at=monday)], [running, undated])
        self.assertEqual([row["id"] for row in self.store.undelivered(7)], [closed, running, undated])

    def test_a_journal_older_than_the_deadline_column_is_upgraded_in_place(self):
        import sqlite3
        with sqlite3.connect(self.store.path) as connection:
            connection.execute("ALTER TABLE leads DROP COLUMN deadline")
        connection.close()
        lead = LeadStore(self.store.path).save(tender_id="t", fingerprint="f", model="m", cost_usd=0, tokens_in=0,
                                               tokens_out=0, tool_calls=0, result=GOOD, card_html="c", deliverable=True,
                                               deadline="2026-10-04T09:00:00Z")
        self.assertEqual(self.store.undelivered(7, open_at="2026-10-05T06:00:00Z"), [])
        self.assertEqual([row["id"] for row in self.store.undelivered(7, open_at="2026-10-03T06:00:00Z")], [lead])

    def test_symlinked_journal_is_rejected(self):
        real = Path(self.tmp.name) / "real.sqlite3"
        real.write_bytes(b"")
        link = Path(self.tmp.name) / "link.sqlite3"
        try:
            link.symlink_to(real)
        except OSError:
            self.skipTest("symlinks unavailable")
        with self.assertRaises(ValueError):
            LeadStore(link)


class RunnerTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.store = LeadStore(self.root / "state" / "leads.sqlite3")

    def source(self, rows):
        path = self.root / "analysis-source.json"
        path.write_text(json.dumps({"rows": rows}, ensure_ascii=False), encoding="utf-8")
        return path

    @staticmethod
    def fake_run(result, cost=0.01, models=None):
        def run(row, *, api_key, model, **kw):
            if models is not None:
                models.append(model)
            return A.AgentRun(tender_id=row["card"]["id"], model=model, final_text=None if result is None else json.dumps(result),
                              cost_usd=cost, tokens_in=10, tokens_out=5, tool_calls=3, tool_text=READ)
        return run

    def test_selects_only_active_new_versions_most_urgent_first(self):
        rows = [make_row("a", deadline="2026-10-09T00:00:00Z"), make_row("b", deadline="2026-10-03T00:00:00Z"),
                make_row("closed", status="ended"), make_row("expired", deadline="2026-09-01T00:00:00Z"),
                make_row("done")]
        self.store.save(tender_id="done", fingerprint=runner.version(rows[-1]), model="m", cost_usd=0, tokens_in=0, tokens_out=0, tool_calls=0,
                        result={"verdict": "reject"}, card_html=None, deliverable=False)
        chosen = runner.select(rows, self.store, include_history=False, now=NOW)
        self.assertEqual([row["card"]["id"] for row in chosen], ["b", "a"])
        self.assertEqual(len(runner.select(rows, self.store, include_history=True, now=NOW)), 4)

    def test_purchase_with_queued_documents_waits_unless_the_deadline_is_close(self):
        queued = [{"title": "ТЗ.pdf", "reason": "run_download_budget"}]
        rows = [dict(make_row("waits", deadline="2026-10-09T00:00:00Z"), gaps=queued),
                dict(make_row("urgent", deadline="2026-10-02T08:00:00Z"), gaps=queued),
                dict(make_row("partial", deadline="2026-10-09T00:00:00Z"), gaps=[{"title": "x", "reason": "per_purchase_file_budget"}])]
        chosen = runner.select(rows, self.store, include_history=False, now=NOW)
        self.assertEqual([row["card"]["id"] for row in chosen], ["urgent", "partial"])

    def test_new_version_of_a_delivered_procurement_is_marked_as_an_update(self):
        runner.process(self.source([make_row()]), self.store, api_key="k", run=self.fake_run(GOOD), now=NOW)
        first = self.store.undelivered(5)[0]
        self.assertNotIn("🔄", first["card_html"])
        self.store.claim(first["id"], 5)
        self.store.record_sent(first["id"], 5, 42)
        runner.process(self.source([make_row(deadline="2026-10-07T09:00:00Z")]), self.store, api_key="k", run=self.fake_run(GOOD), now=NOW)
        update = self.store.undelivered(5)[0]["card_html"]
        self.assertIn("🔄 <i>Закупка изменилась", update.splitlines()[1])
        self.assertEqual(update.count("<i>"), update.count("</i>"))

    def test_only_what_the_agent_reads_makes_a_new_version(self):
        runner.process(self.source([make_row()]), self.store, api_key="k", run=self.fake_run(GOOD), now=NOW)
        reread = dict(make_row(), modified="2026-10-02T00:00:00Z", gaps=[{"title": "Схема.dwg", "reason": "unsupported_format"}])
        reread["card"]["document_state"] = "retrieved_partial"
        self.assertEqual(runner.select([reread], self.store, include_history=False, now=NOW), [])
        for change in (lambda row: row["card"].update(description="Нужны 4 сервера"),
                       lambda row: row["card"]["documents"][0].update(text="ТЗ, редакция 2"),
                       lambda row: row["card"]["documents"].append({"id": "d2", "title": "Спецификация", "text": "2 шт."}),
                       lambda row: row["opportunity"].update(acceptance_end_date="2026-10-07T09:00:00Z")):
            changed = make_row()
            change(changed)
            self.assertEqual(len(runner.select([changed], self.store, include_history=False, now=NOW)), 1)

    def test_grade_a_result_is_grounded_stored_and_not_repeated(self):
        models = []
        counts = runner.process(self.source([make_row()]), self.store, api_key="k", run=self.fake_run(GOOD, models=models), now=NOW)
        self.assertEqual((counts["saved"], counts["deliverable"], counts["retried_with_fallback"]), (1, 1, 0))
        self.assertEqual(models, [A.DEFAULT_MODEL])
        self.assertIn("☎", self.store.undelivered(5)[0]["card_html"])
        again = runner.process(self.source([make_row()]), self.store, api_key="k", run=self.fake_run(GOOD), now=NOW)
        self.assertEqual((again["selected"], again["saved"]), (0, 0))

    def test_weaker_grade_retries_once_with_fallback_and_keeps_the_better_result(self):
        weak = json.loads(json.dumps(GOOD))
        weak["contacts"][0]["name"] = None  # phone only -> grade B
        models = []

        def run(row, *, api_key, model, **kw):
            payload = weak if model == A.DEFAULT_MODEL else GOOD
            models.append(model)
            return A.AgentRun(tender_id=row["card"]["id"], model=model, final_text=json.dumps(payload), cost_usd=0.01,
                              tokens_in=1, tokens_out=1, tool_calls=1, tool_text=READ)
        counts = runner.process(self.source([make_row()]), self.store, api_key="k", run=run, now=NOW)
        self.assertEqual(models, [A.DEFAULT_MODEL, A.FALLBACK_MODEL])
        self.assertEqual(counts["retried_with_fallback"], 1)
        self.assertEqual(self.store.get(1)["grade"], "A")
        self.assertEqual(counts["cost_usd"], 0.02)

    def test_reject_by_the_cheap_model_gets_a_second_opinion(self):
        reject = {"verdict": "reject", "grade": None, "one_line": "единичная закупка"}
        models = []

        def run(row, *, api_key, model, **kw):
            models.append(model)
            payload = reject if model == A.DEFAULT_MODEL else GOOD
            return A.AgentRun(tender_id=row["card"]["id"], model=model, final_text=json.dumps(payload), cost_usd=0.01,
                              tokens_in=1, tokens_out=1, tool_calls=1, tool_text=READ)
        counts = runner.process(self.source([make_row()]), self.store, api_key="k", run=run, now=NOW)
        self.assertEqual(models, [A.DEFAULT_MODEL, A.FALLBACK_MODEL])
        self.assertEqual((counts["deliverable"], self.store.get(1)["verdict"]), (1, "lead"))

    def test_unparsable_results_are_retried_once_then_recorded_as_failed(self):
        source = self.source([make_row()])
        first = runner.process(source, self.store, api_key="k", run=self.fake_run(None), fallback_model=None, now=NOW)
        self.assertEqual((first["failed"], first["saved"]), (1, 0))
        second = runner.process(source, self.store, api_key="k", run=self.fake_run(None), fallback_model=None, now=NOW)
        self.assertEqual((second["failed"], second["saved"]), (0, 1))
        self.assertEqual(self.store.get(1)["verdict"], "failed")
        self.assertEqual(self.store.undelivered(5), [])

    def test_reject_and_unverifiable_contacts_are_never_delivered(self):
        counts = runner.process(self.source([make_row("r")]), self.store, api_key="k", fallback_model=None,
                                run=self.fake_run({"verdict": "reject", "grade": None, "one_line": "софт"}), now=NOW)
        self.assertEqual((counts["saved"], counts["deliverable"]), (1, 0))

    def test_run_cost_cap_stops_further_investigations(self):
        rows = [make_row(f"t{i}") for i in range(4)]
        counts = runner.process(self.source(rows), self.store, api_key="k", run=self.fake_run(GOOD, cost=0.6),
                                max_run_cost=1.0, fallback_model=None, now=NOW)
        self.assertEqual(counts["saved"], 2)

    def test_run_outcome_is_left_for_the_health_check(self):
        state = self.root / "state"
        with patch.object(runner, "process", return_value={"selected": 3, "saved": 2, "failed": 1}), \
                patch.dict("os.environ", {"OPENROUTER_API_KEY": "k"}), patch("builtins.print"):
            self.assertEqual(runner.main(["--source", "x", "--state-root", str(state)]), 0)
        beat = json.loads((state / "leads-research.json").read_text(encoding="utf-8"))
        self.assertEqual((beat["selected"], beat["saved"]), (3, 2))
        self.assertIn("at", beat)
        with patch.object(runner, "process", side_effect=RuntimeError("proxy down")), \
                patch.dict("os.environ", {"OPENROUTER_API_KEY": "k"}), self.assertRaises(RuntimeError):
            runner.main(["--source", "x", "--state-root", str(state)])
        self.assertEqual(json.loads((state / "leads-research.json").read_text(encoding="utf-8"))["error"], "RuntimeError")


class HealthTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.bundles, self.state = Path(self.tmp.name) / "bundles", Path(self.tmp.name) / "state"
        self.bundles.mkdir()
        self.state.mkdir()

    def write(self, root, name, payload):
        (root / name).write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")

    def test_fresh_source_and_research_report_nothing(self):
        self.write(self.bundles, "source-refresh.json", {"collected_at": "2026-10-01T08:30:00Z"})
        self.write(self.state, "leads-research.json", {"at": "2026-10-01T07:30:00Z", "selected": 2, "saved": 2, "failed": 0})
        self.assertEqual(health.problems(self.bundles, self.state, NOW), {})
        self.assertEqual(health.plan({}, {}, {}, NOW), ([], {"problems": {}, "quarantined": []}))

    def test_stale_source_is_reported_once_reminded_daily_and_closed_on_recovery(self):
        self.write(self.bundles, "source-refresh.json", {"collected_at": "2026-09-30T08:27:00Z"})
        found = health.problems(self.bundles, self.state, NOW,
                                source_error="SourceBoundaryError: invalid purchase description https://x.test/secret-path")
        self.assertIn("не обновляется 24 ч", found["source"])
        self.assertIn("invalid purchase description", found["source"])
        self.assertNotIn("secret-path", found["source"])
        messages, state = health.plan({}, found, {}, NOW)
        self.assertEqual(len(messages), 1)
        self.assertEqual(health.plan(state, found, {}, NOW.replace(hour=15))[0], [])
        reminder, _ = health.plan(state, found, {}, NOW.replace(day=2, hour=10))
        self.assertIn("не устранена", reminder[0])
        recovered, after = health.plan(state, {}, {}, NOW.replace(hour=15))
        self.assertIn("восстановлен", recovered[0])
        self.assertEqual(after["problems"], {})

    def test_missing_source_failed_research_and_stuck_delivery_are_problems(self):
        self.write(self.state, "leads-research.json", {"at": "2026-10-01T07:30:00Z", "selected": 4, "saved": 0, "failed": 4})
        store = LeadStore(self.state / "leads.sqlite3")
        lead = store.save(tender_id="t", fingerprint="f", model="m", cost_usd=0, tokens_in=0, tokens_out=0, tool_calls=0,
                          result=GOOD, card_html="<b>card</b>", deliverable=True)
        store.claim(lead, 5)
        store.record_uncertain(lead, 5)
        found = health.problems(self.bundles, self.state, NOW)
        self.assertEqual(set(found), {"source", "research", "delivery"})
        self.write(self.state, "leads-research.json", {"at": "2026-10-01T07:30:00Z", "error": "ConnectError"})
        self.assertIn("ConnectError", health.problems(self.bundles, self.state, NOW)["research"])
        self.write(self.state, "leads-research.json", {"at": "2026-09-29T07:30:00Z", "selected": 0})
        self.assertIn("не запускалось", health.problems(self.bundles, self.state, NOW)["research"])

    def test_quarantined_purchase_is_reported_once(self):
        self.write(self.bundles, "source-quarantine.json", {"items": [{"id": "example-source:101-002", "reason": "invalid file inventory"}]})
        skipped = health.quarantined(self.bundles)
        messages, state = health.plan({}, {}, skipped, NOW)
        self.assertIn("№ 101-002: invalid file inventory", messages[0])
        self.assertEqual(health.plan(state, {}, skipped, NOW)[0], [])

    def test_empty_slot_line_counts_waiting_purchases_and_is_silent_after_a_delivery(self):
        rows = [make_row("a"), make_row("b"), make_row("old", status="ended")]
        self.write(self.bundles, "analysis-source.json", {"collected_at": "2026-10-01T08:30:00Z", "rows": rows})
        store = LeadStore(self.state / "leads.sqlite3")
        lead = store.save(tender_id="a", fingerprint=runner.version(rows[0]), model="m", cost_usd=0, tokens_in=0, tokens_out=0,
                          tool_calls=0, result=GOOD, card_html="<b>card</b>", deliverable=True)
        line = health.slot_line(self.bundles, self.state, 5, NOW)
        self.assertIn("Активных профильных закупок: 2, ждут исследования: 1", line)
        self.assertIn("01.10 11:30 МСК", line)
        store.save(tender_id="b", fingerprint=runner.version(rows[1]), model="m", cost_usd=0, tokens_in=0, tokens_out=0,
                   tool_calls=0, result={"verdict": "reject"}, card_html=None, deliverable=False)
        self.assertIn("все разобраны", health.slot_line(self.bundles, self.state, 5, NOW))
        store.claim(lead, 5)
        store.record_sent(lead, 5, 77)
        self.assertIsNone(health.slot_line(self.bundles, self.state, 5, datetime.now(timezone.utc)))


class ExportTest(unittest.TestCase):
    NOW = datetime(2026, 10, 1, 12, 30, tzinfo=export.MSK)

    def item(self, deadline, name="Иванов Иван", taken="2026-10-01T09:30:00Z", grade="A", **extra):
        return {"id": 1, "tender_id": "t", "grade": grade, "taken_at": taken,
                "card_html": '<a href="https://tenders.example.org/process/x">Закупка</a>',
                "result": {"customer": {"name": "ООО «Учебный заказчик»"}, "signal": {"what_they_buy": "2 сервера <b>&</b>", "deadline": deadline},
                           "contacts": [{"name": name, "role": "закупки", "phone": "+7 (000) 000-00-00", "email": "i@example.org"},
                                        {"name": None, "role": "приёмная", "phone": "8 000 000-00-99"}],
                           "talk_track": {"hooks": ["один", "два"]}, **extra}}

    def test_deadlines_are_parsed_in_dotted_and_iso_forms(self):
        msk = export.MSK
        self.assertEqual(export.parse_deadline("до 01.10.2026 09:00 МСК"), (datetime(2026, 10, 1, 9, 0, tzinfo=msk), True))
        self.assertEqual(export.parse_deadline("02.10.2026"), (datetime(2026, 10, 2, 23, 59, tzinfo=msk), False))
        self.assertEqual(export.parse_deadline("2026-10-02T11:00:00Z"), (datetime(2026, 10, 2, 14, 0, tzinfo=msk), True))
        self.assertEqual(export.parse_deadline("2026-10-01 09:00"), (datetime(2026, 10, 1, 9, 0, tzinfo=msk), True))
        for bad in ("через 7 рабочих дней", "", None, "31.02.2026"):
            self.assertIsNone(export.parse_deadline(bad))

    def test_workbook_is_a_passive_xlsx_formatted_for_a_caller(self):
        import io
        import zipfile
        payload = export.build_workbook([self.item("02.10.2026 09:00")], now=self.NOW)
        archive = zipfile.ZipFile(io.BytesIO(payload))
        self.assertTrue(all(not name.endswith((".bin", ".vba")) for name in archive.namelist()))
        self.assertNotIn("<f>", archive.read("xl/worksheets/sheet1.xml").decode("utf-8"))  # no formulas
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl is not installed")
        sheet = openpyxl.load_workbook(io.BytesIO(payload)).active
        self.assertEqual(sheet["A1"].value, "Лиды в работе")
        self.assertIn("лидов: 1", sheet["A2"].value)
        header = [cell.value for cell in sheet[4]]
        self.assertEqual(header[:5], ["№", "Заказчик", "Контакт", "Телефон", "Срок приёма"])
        row = dict(zip(header, [cell.value for cell in sheet[5]]))
        self.assertEqual((row["Контакт"], row["Телефон"], row["Оценка"]), ("Иванов Иван", "+7 (000) 000-00-00", "A"))
        self.assertEqual(row["Срок приёма"], datetime(2026, 10, 2, 9, 0))  # a real date, sortable in Excel
        self.assertEqual(row["Взято в работу"], datetime(2026, 10, 1, 12, 30))
        self.assertEqual(row["Что покупают"], "2 сервера <b>&</b>")
        self.assertEqual(row["Другие контакты"], "приёмная — 8 000 000-00-99")
        self.assertEqual(sheet.cell(5, len(header)).hyperlink.target, "https://tenders.example.org/process/x")
        self.assertEqual(sheet.freeze_panes, "E5")
        self.assertEqual(sheet.auto_filter.ref, "A4:N5")
        self.assertEqual(sheet.cell(5, header.index("Оценка") + 1).fill.fgColor.rgb, "FFD3F0DD")  # grade A is green

    def test_rows_sort_soonest_first_undated_next_expired_last_and_urgent_is_highlighted(self):
        import io
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl is not installed")
        items = [self.item("30.09.2026 10:00", name="Просрочен"), self.item("через неделю", name="Без даты"),
                 self.item("10.10.2026 10:00", name="Поздно"), self.item("02.10.2026 09:00", name="Срочно")]
        sheet = openpyxl.load_workbook(io.BytesIO(export.build_workbook(items, now=self.NOW))).active
        self.assertEqual([sheet.cell(r, 3).value for r in range(5, 9)], ["Срочно", "Поздно", "Без даты", "Просрочен"])
        self.assertEqual(sheet["E5"].fill.fgColor.rgb, "FFFDE4E1")      # within 48 hours
        self.assertNotEqual(sheet["E6"].fill.fgColor.rgb, "FFFDE4E1")   # a later deadline is not flagged
        self.assertIn("срочно", sheet["A2"].value)
        self.assertEqual(sheet["E7"].value, "через неделю")

    def test_new_leads_workbook_has_its_own_title_and_date_header(self):
        import io
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl is not installed")
        item = self.item("02.10.2026 09:00")
        item["stamp"] = "2026-10-01T06:00:00Z"
        sheet = openpyxl.load_workbook(io.BytesIO(export.build_workbook([item], kind="new", now=self.NOW))).active
        self.assertEqual((sheet.title, sheet["A1"].value), ("Новые лиды", "Новые лиды"))
        self.assertIn("ещё не нажимали", sheet["A2"].value)
        self.assertEqual(sheet["M4"].value, "Получен")
        self.assertEqual(sheet["M5"].value, datetime(2026, 10, 1, 9, 0))  # delivery time in Moscow

    def test_store_new_leads_are_delivered_cards_without_any_answer_from_that_user(self):
        with tempfile.TemporaryDirectory() as directory:
            store = LeadStore(Path(directory) / "leads.sqlite3")
            ids = [store.save(tender_id=f"t{i}", fingerprint="f", model="m", cost_usd=0, tokens_in=0, tokens_out=0,
                              tool_calls=0, result=GOOD, card_html="c", deliverable=True) for i in range(4)]
            for lead in ids[:3]:
                store.claim(lead, 5)
                store.record_sent(lead, 5, 100 + lead)
            store.claim(ids[3], 5)  # claimed but never confirmed: not "new" for the manager
            store.add_feedback(ids[1], 5, 5, "skip")
            store.add_feedback(ids[2], 5, 6, "work")  # another user's answer does not hide it from user 5
            self.assertEqual([item["id"] for item in store.new_leads(5, 5)], [ids[0], ids[2]])
            self.assertEqual([item["id"] for item in store.new_leads(5, 6)], [ids[0], ids[1]])
            self.assertEqual(store.new_leads(7, 7), [])

    def test_comment_column_shows_the_managers_notes_in_a_sticky_note_colour(self):
        import io
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl is not installed")
        item = self.item("02.10.2026 09:00")
        item["comments"] = [{"at": "2026-10-01T07:05:00Z", "text": "позвонил, секретарь переключила на ИТ"},
                            {"at": "2026-10-01T11:30:00Z", "text": "перезвонить во вторник"}]
        sheet = openpyxl.load_workbook(io.BytesIO(export.build_workbook([item, self.item("05.10.2026")], now=self.NOW))).active
        header = [cell.value for cell in sheet[4]]
        column = header.index("Комментарий") + 1
        self.assertEqual(header[:6], ["№", "Заказчик", "Контакт", "Телефон", "Срок приёма", "Комментарий"])
        self.assertEqual(sheet.cell(5, column).value,
                         "01.10 10:05 — позвонил, секретарь переключила на ИТ\n01.10 14:30 — перезвонить во вторник")
        self.assertIsNone(sheet.cell(6, column).value)  # a lead without notes keeps an empty note cell
        self.assertEqual(sheet.cell(6, column).fill.fgColor.rgb, "FFFFF6D5")

    def test_empty_unsafe_links_and_control_characters_are_handled(self):
        import io
        import zipfile
        item = self.item("02.10.2026")
        item["card_html"] = '<a href="http://insecure.example/x">x</a>'
        item["result"]["customer"]["name"] = "Бо" + chr(0) + "гус"
        archive = zipfile.ZipFile(io.BytesIO(export.build_workbook([item], now=self.NOW)))
        self.assertNotIn("xl/worksheets/_rels/sheet1.xml.rels", archive.namelist())
        self.assertNotIn(chr(0), archive.read("xl/worksheets/sheet1.xml").decode("utf-8"))
        self.assertTrue(export.build_workbook([], now=self.NOW))

    def test_store_returns_only_latest_work_verdicts_of_that_user(self):
        with tempfile.TemporaryDirectory() as directory:
            store = LeadStore(Path(directory) / "leads.sqlite3")
            ids = [store.save(tender_id=f"t{i}", fingerprint="f", model="m", cost_usd=0, tokens_in=0, tokens_out=0,
                              tool_calls=0, result=GOOD, card_html="c", deliverable=True) for i in range(3)]
            for lead in ids:
                store.claim(lead, 5)
            store.add_feedback(ids[0], 5, 5, "work")
            store.add_feedback(ids[1], 5, 5, "work")
            store.add_feedback(ids[1], 5, 5, "skip")
            store.add_feedback(ids[2], 5, 6, "work")  # another user's verdict never leaks into this export
            self.assertEqual([item["id"] for item in store.in_work(5)], [ids[0]])
            self.assertEqual([item["id"] for item in store.in_work(6)], [ids[2]])


class CommentStoreTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.store = LeadStore(Path(self.tmp.name) / "leads.sqlite3")
        self.lead = self.store.save(tender_id="t", fingerprint="f", model="m", cost_usd=0, tokens_in=0, tokens_out=0,
                                    tool_calls=0, result=GOOD, card_html="c", deliverable=True)
        self.store.claim(self.lead, 5)
        self.store.record_sent(self.lead, 5, 77)

    def test_only_a_confirmed_delivered_card_accepts_a_note(self):
        self.assertTrue(self.store.has_card(5, 77))
        self.assertFalse(self.store.has_card(5, 78))
        self.assertFalse(self.store.has_card(6, 77))
        self.assertFalse(self.store.add_comment(5, 5, 78, "не карточка"))
        self.assertFalse(self.store.add_comment(6, 6, 77, "чужой чат"))
        self.assertFalse(self.store.add_comment(5, 5, 77, "   \n  "))
        self.assertTrue(self.store.add_comment(5, 5, 77, "позвонил,   секретарь\n\n переключила"))

    def test_notes_travel_with_the_lead_into_both_exports_for_that_user_only(self):
        self.store.add_comment(5, 5, 77, "первая")
        self.store.add_comment(5, 5, 77, "вторая")
        self.assertEqual([c["text"] for c in self.store.new_leads(5, 5)[0]["comments"]], ["первая", "вторая"])
        self.assertEqual(self.store.new_leads(5, 6)[0]["comments"], [])  # another manager does not see these notes
        self.store.add_feedback(self.lead, 5, 5, "work")
        self.assertEqual([c["text"] for c in self.store.in_work(5)[0]["comments"]], ["первая", "вторая"])
        self.assertEqual(self.store.new_leads(5, 5), [])

    def test_long_notes_are_capped(self):
        self.store.add_comment(5, 5, 77, "я" * 5000)
        self.assertEqual(len(self.store.new_leads(5, 5)[0]["comments"][0]["text"]), 2000)


class DeliveryAndFeedbackTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.store = LeadStore(Path(self.tmp.name) / "leads.sqlite3")
        self.lead = self.store.save(tender_id="t", fingerprint="f", model="m", cost_usd=0, tokens_in=0, tokens_out=0,
                                    tool_calls=0, result=GOOD, card_html="<b>card</b>", deliverable=True)

    def deliver(self, bot, recipients=(5,)):
        return asyncio.run(delivery.deliver(self.store, bot, list(recipients), pause=0, keyboard=lambda lead_id: f"kb{lead_id}"))

    def test_sends_once_with_keyboard_and_records_receipt(self):
        bot = SimpleNamespace(send_message=AsyncMock(return_value=SimpleNamespace(message_id=42)))
        self.assertEqual(self.deliver(bot)["sent"], 1)
        kwargs = bot.send_message.await_args.kwargs
        self.assertEqual((kwargs["chat_id"], kwargs["parse_mode"], kwargs["reply_markup"]), (5, "HTML", f"kb{self.lead}"))
        self.assertEqual(self.deliver(bot)["sent"], 0)
        self.assertEqual(bot.send_message.await_count, 1)

    def test_ambiguous_failure_is_not_retried(self):
        bot = SimpleNamespace(send_message=AsyncMock(side_effect=TimeoutError()))
        self.assertEqual(self.deliver(bot)["uncertain"], 1)
        self.assertEqual(self.deliver(SimpleNamespace(send_message=AsyncMock()))["sent"], 0)

    def test_flood_control_releases_claim_and_stops(self):
        class RetryAfter(Exception):
            retry_after = 30
        bot = SimpleNamespace(send_message=AsyncMock(side_effect=RetryAfter()))
        self.assertEqual(self.deliver(bot)["rate_limited"], 1)
        ok = SimpleNamespace(send_message=AsyncMock(return_value=SimpleNamespace(message_id=1)))
        self.assertEqual(self.deliver(ok)["sent"], 1)

    def test_callback_data_roundtrip_and_layout(self):
        self.assertEqual(feedback.parse("lf:12:skip"), (12, "skip"))
        for bad in ("lf:0:work", "lf:1:call", "lf:1:rm", "x:1:call", "lf:1:call:x", "", None):
            self.assertIsNone(feedback.parse(bad))
        rows = feedback.button_rows(7, {"work"})
        self.assertEqual([len(row) for row in rows], [2])
        self.assertTrue(rows[0][0][0].startswith("✓"))
        self.assertFalse(rows[0][1][0].startswith("✓"))
        self.assertTrue(all(len(data.encode()) <= 64 for row in rows for _, data in row))

    def callback(self, data, user_id=5, chat_id=5, chat_type="private", allowed=True):
        query = SimpleNamespace(data=data, from_user=SimpleNamespace(id=user_id), message=SimpleNamespace(chat=SimpleNamespace(id=chat_id, type=chat_type)),
                                answer=AsyncMock(), edit_message_reply_markup=AsyncMock())
        with patch.object(feedback, "_allowed", return_value=allowed), patch.object(feedback, "keyboard", side_effect=lambda lead_id, done=(): ("kb", tuple(done))):
            asyncio.run(feedback.on_callback(SimpleNamespace(callback_query=query), None, store=self.store))
        return query

    def test_button_press_is_recorded_for_delivered_lead_only_by_allowed_private_user(self):
        self.store.claim(self.lead, 5)
        query = self.callback(f"lf:{self.lead}:work")
        self.assertEqual(self.store.feedback(self.lead), ["work"])
        query.answer.assert_awaited_with("Записано")
        self.assertEqual(query.edit_message_reply_markup.await_args.kwargs["reply_markup"], ("kb", ("work",)))
        for kwargs in ({"allowed": False}, {"chat_type": "group"}, {"chat_id": 6}, {"user_id": 6}):
            self.callback(f"lf:{self.lead}:skip", **kwargs)
        corrected = self.callback(f"lf:{self.lead}:skip")  # a correction: the check mark moves to the last verdict
        self.assertEqual(self.store.feedback(self.lead), ["work", "skip"])
        self.assertEqual(corrected.edit_message_reply_markup.await_args.kwargs["reply_markup"], ("kb", ("skip",)))
        self.callback("lf:999:work")
        self.assertEqual(self.store.feedback(self.lead), ["work", "skip"])


if __name__ == "__main__":
    unittest.main()
