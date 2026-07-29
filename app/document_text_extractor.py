from __future__ import annotations

import argparse
import logging
import shutil
import subprocess
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

import fitz
import psycopg
import xlrd
from docx import Document
from openpyxl import load_workbook
from psycopg.rows import dict_row

from app.config import settings
from app.document_archive_utils import (
    MAX_ARCHIVE_DEPTH,
    ArchiveMemberCandidate,
    SafeArchiveMember,
    is_safe_child_path,
    select_safe_archive_members,
)


APP_DIR = Path("/app")

logger = logging.getLogger(__name__)

RAR_BACKEND_CANDIDATES = ("unar", "unrar", "bsdtar", "7z", "7zz")
LEGACY_DOC_TOOL = "antiword"
LEGACY_DOC_TIMEOUT_SECONDS = 60
LEGACY_DOC_ERROR_PREVIEW_CHARS = 500
LEGACY_DOC_COMMAND_ARGS = (
    ("utf8_mapping", ("-m", "UTF-8.txt")),
    ("default_mapping", ()),
)

DOCX_TEXT_XML_NAMES = {
    "word/document.xml",
    "word/footnotes.xml",
    "word/endnotes.xml",
    "word/comments.xml",
    "word/glossary/document.xml",
}

_EMPTY_EXTRACTION_REASONS: dict[str, str] = {}


def clean_text(text: str) -> str:
    lines = []
    for line in text.splitlines():
        line = " ".join(line.strip().split())
        if line:
            lines.append(line)
    return "\n".join(lines)


def clean_inline_text(text: str) -> str:
    return " ".join(str(text).strip().split())


def remember_empty_extraction_reason(path: Path, reason: str) -> None:
    _EMPTY_EXTRACTION_REASONS[str(path)] = reason


def remembered_empty_extraction_reason(path: Path) -> str | None:
    return _EMPTY_EXTRACTION_REASONS.get(str(path))


def docx_container_text_parts(container: Any) -> list[str]:
    parts = []

    for paragraph in getattr(container, "paragraphs", []) or []:
        text = clean_inline_text(getattr(paragraph, "text", "") or "")
        if text:
            parts.append(text)

    for table in getattr(container, "tables", []) or []:
        parts.extend(docx_table_text_parts(table))

    return parts


def docx_table_text_parts(table: Any) -> list[str]:
    parts = []

    for row in getattr(table, "rows", []) or []:
        cells = []
        for cell in getattr(row, "cells", []) or []:
            cell_text = clean_inline_text(" ".join(docx_container_text_parts(cell)))
            if not cell_text:
                cell_text = clean_inline_text(getattr(cell, "text", "") or "")
            if cell_text:
                cells.append(cell_text)

        if cells:
            parts.append(" | ".join(cells))

    return parts


def docx_section_text_parts(doc: Any) -> list[str]:
    parts = []

    for section in getattr(doc, "sections", []) or []:
        for attr in (
            "header",
            "first_page_header",
            "even_page_header",
            "footer",
            "first_page_footer",
            "even_page_footer",
        ):
            container = getattr(section, attr, None)
            if container is not None:
                parts.extend(docx_container_text_parts(container))

    return parts


def xml_local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def is_docx_text_xml_name(name: str) -> bool:
    normalized = name.replace("\\", "/")
    if normalized in DOCX_TEXT_XML_NAMES:
        return True
    if normalized.startswith("word/header") and normalized.endswith(".xml"):
        return True
    if normalized.startswith("word/footer") and normalized.endswith(".xml"):
        return True
    return False


def docx_xml_node_text(node: ET.Element) -> str:
    fragments = []

    for element in node.iter():
        local_name = xml_local_name(str(element.tag))
        if local_name in {"t", "instrText"} and element.text:
            fragments.append(element.text)
        elif local_name == "tab":
            fragments.append("\t")
        elif local_name in {"br", "cr"}:
            fragments.append("\n")

    return clean_inline_text("".join(fragments))


def docx_xml_text_lines(xml_bytes: bytes) -> list[str]:
    root = ET.fromstring(xml_bytes)
    lines = []

    for paragraph in root.iter():
        if xml_local_name(str(paragraph.tag)) != "p":
            continue
        text = docx_xml_node_text(paragraph)
        if text:
            lines.append(text)

    if lines:
        return lines

    fallback = docx_xml_node_text(root)
    if fallback:
        return [fallback]

    return []


def extract_docx_xml_text(path: Path) -> str:
    parts = []

    try:
        with zipfile.ZipFile(path) as archive:
            names = sorted(
                name for name in archive.namelist() if is_docx_text_xml_name(name)
            )
            for name in names:
                try:
                    parts.extend(docx_xml_text_lines(archive.read(name)))
                except ET.ParseError as exc:
                    logger.warning(
                        "DOCX XML text extraction skipped member: path=%s member=%s reason=%s",
                        path,
                        name,
                        exc,
                    )
    except (OSError, zipfile.BadZipFile) as exc:
        logger.warning("DOCX XML text extraction failed: path=%s reason=%s", path, exc)
        return ""

    return clean_text("\n".join(parts))


def extract_docx(path: Path) -> str:
    parts = []
    primary_error: Exception | None = None

    try:
        doc = Document(path)
        parts.extend(docx_container_text_parts(doc))
        parts.extend(docx_section_text_parts(doc))
    except Exception as exc:
        primary_error = exc
        logger.warning(
            "DOCX primary text extraction failed: path=%s reason=%s",
            path,
            exc,
        )

    text = clean_text("\n".join(parts))
    if text:
        return text

    text = extract_docx_xml_text(path)
    if text:
        return text

    remember_empty_extraction_reason(path, "docx_empty_text")
    if primary_error is not None:
        raise primary_error
    return ""


def format_spreadsheet_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat(timespec="seconds")
    return clean_inline_text(str(value))


def close_workbook(wb: Any) -> None:
    close = getattr(wb, "close", None)
    if callable(close):
        close()


def extract_xlsx_workbook_text(path: Path, *, data_only: bool) -> str:
    wb = load_workbook(path, read_only=True, data_only=data_only)
    parts = []

    try:
        for sheet in wb.worksheets:
            sheet_parts = []

            for row in sheet.iter_rows():
                values = []
                for cell in row:
                    text = format_spreadsheet_value(getattr(cell, "value", None))
                    if not text:
                        continue

                    coordinate = clean_inline_text(getattr(cell, "coordinate", "") or "")
                    values.append(f"{coordinate}: {text}" if coordinate else text)

                if values:
                    sheet_parts.append(" | ".join(values))

            if sheet_parts:
                parts.append(f"Sheet: {sheet.title}")
                parts.extend(sheet_parts)
    finally:
        close_workbook(wb)

    return clean_text("\n".join(parts))


def extract_xlsx(path: Path) -> str:
    text = extract_xlsx_workbook_text(path, data_only=True)
    if text:
        return text

    text = extract_xlsx_workbook_text(path, data_only=False)
    if text:
        return text

    remember_empty_extraction_reason(path, "xlsx_empty_workbook")
    return ""


def extract_xls(path: Path) -> str:
    wb = xlrd.open_workbook(path)
    parts = []

    for sheet in wb.sheets():
        parts.append(f"Sheet: {sheet.name}")

        for row_idx in range(sheet.nrows):
            values = []
            for col_idx in range(sheet.ncols):
                value = sheet.cell_value(row_idx, col_idx)
                if value is not None and str(value).strip():
                    values.append(str(value).strip())
            if values:
                parts.append(" | ".join(values))

    return clean_text("\n".join(parts))


def legacy_doc_warning(path: Path, reason: str, detail: str | None = None) -> str:
    remember_empty_extraction_reason(path, reason)

    if detail:
        logger.warning(
            "Legacy DOC extraction failed: path=%s reason=%s detail=%s",
            path,
            reason,
            detail[:LEGACY_DOC_ERROR_PREVIEW_CHARS],
        )
    else:
        logger.warning("Legacy DOC extraction failed: path=%s reason=%s", path, reason)

    return ""


def decode_legacy_doc_output(value: bytes | str | None) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value

    for encoding in ("utf-8", "cp1251", "latin-1"):
        try:
            return value.decode(encoding)
        except UnicodeDecodeError:
            continue

    return value.decode("utf-8", errors="replace")


def extract_doc_legacy(path: Path) -> str:
    if not path.exists():
        return legacy_doc_warning(path, "file_not_found")

    tool_path = shutil.which(LEGACY_DOC_TOOL)
    if not tool_path:
        return legacy_doc_warning(path, f"{LEGACY_DOC_TOOL}_not_found")

    last_reason = f"{LEGACY_DOC_TOOL}_empty_output"
    last_stderr = ""

    for attempt_name, command_args in LEGACY_DOC_COMMAND_ARGS:
        try:
            result = subprocess.run(
                [tool_path, *command_args, str(path)],
                capture_output=True,
                check=False,
                timeout=LEGACY_DOC_TIMEOUT_SECONDS,
            )
        except subprocess.TimeoutExpired:
            return legacy_doc_warning(
                path,
                f"{LEGACY_DOC_TOOL}_timeout",
                f"attempt={attempt_name}",
            )
        except OSError as exc:
            return legacy_doc_warning(
                path,
                f"{LEGACY_DOC_TOOL}_failed_to_start",
                str(exc),
            )

        text = clean_text(decode_legacy_doc_output(result.stdout))
        stderr = clean_text(decode_legacy_doc_output(result.stderr))

        if result.returncode != 0:
            reason = f"{LEGACY_DOC_TOOL}_exit_{result.returncode}"
            if text:
                logger.warning(
                    "Legacy DOC extraction returned text with warning: "
                    "path=%s attempt=%s reason=%s detail=%s",
                    path,
                    attempt_name,
                    reason,
                    stderr[:LEGACY_DOC_ERROR_PREVIEW_CHARS],
                )
                return text

            logger.warning(
                "Legacy DOC extraction attempt failed: "
                "path=%s attempt=%s reason=%s detail=%s",
                path,
                attempt_name,
                reason,
                stderr[:LEGACY_DOC_ERROR_PREVIEW_CHARS],
            )
            last_reason = reason
            last_stderr = stderr
            continue

        if text:
            return text

        logger.warning(
            "Legacy DOC extraction attempt returned empty text: "
            "path=%s attempt=%s detail=%s",
            path,
            attempt_name,
            stderr[:LEGACY_DOC_ERROR_PREVIEW_CHARS],
        )
        last_reason = f"{LEGACY_DOC_TOOL}_empty_output"
        last_stderr = stderr

    return legacy_doc_warning(path, last_reason, last_stderr)


def extract_pdf(path: Path) -> str:
    parts = []

    with fitz.open(path) as doc:
        for page_number, page in enumerate(doc, start=1):
            text = page.get_text("text")
            if text.strip():
                parts.append(f"Page {page_number}")
                parts.append(text)

    return clean_text("\n".join(parts))


def archive_warning(path: Path, reason: str) -> str:
    logger.warning("Archive extraction skipped: archive=%s reason=%s", path, reason)
    return ""


def archive_member_warning(path: Path, member_name: str, reason: str) -> str:
    logger.warning(
        "Archive member extraction skipped: archive=%s member=%r reason=%s",
        path,
        member_name,
        reason,
    )
    return ""


def copy_zip_member_safely(
    archive: zipfile.ZipFile,
    member: SafeArchiveMember,
    target_root: Path,
) -> Path:
    target_path = target_root / member.safe_path
    if not is_safe_child_path(target_root, target_path):
        raise ValueError("unsafe target path")

    target_path.parent.mkdir(parents=True, exist_ok=True)

    with archive.open(member.name) as source, target_path.open("wb") as target:
        shutil.copyfileobj(source, target, length=1024 * 1024)

    return target_path


def extracted_member_paths(
    target_root: Path,
    members: list[SafeArchiveMember],
    archive_path: Path,
) -> list[tuple[Path, str]]:
    paths = []

    for member in members:
        child = target_root / member.safe_path
        if not is_safe_child_path(target_root, child):
            logger.warning(
                "Archive member escaped target dir: archive=%s member=%r target=%s",
                archive_path,
                member.name,
                child,
            )
            continue

        if child.is_symlink():
            logger.warning(
                "Archive member symlink skipped: archive=%s member=%r target=%s",
                archive_path,
                member.name,
                child,
            )
            continue

        if child.is_file():
            paths.append((child, str(member.safe_path)))

    return paths


def extract_archive_member_texts(
    archive_path: Path,
    member_paths: list[tuple[Path, str]],
    archive_depth: int,
) -> str:
    parts = []

    for child, member_name in member_paths:
        try:
            extracted = extract_file(child, archive_depth=archive_depth)
        except Exception as exc:
            extracted = archive_member_warning(archive_path, member_name, str(exc))

        if extracted:
            parts.append(f"\n\n===== File inside archive: {member_name} =====\n")
            parts.append(extracted)

    return clean_text("\n".join(parts))


def extract_zip(path: Path, archive_depth: int = 1) -> str:
    parts = []

    try:
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp_path = Path(tmpdir)

            with zipfile.ZipFile(path) as archive:
                infos = archive.infolist()
                encrypted_names = {
                    info.filename for info in infos if info.flag_bits & 0x1
                }
                candidates = [
                    ArchiveMemberCandidate(
                        name=info.filename,
                        size=info.file_size,
                        is_dir=info.is_dir(),
                    )
                    for info in infos
                ]
                members = select_safe_archive_members(
                    candidates,
                    archive_path=path,
                    logger=logger,
                )

                for member in members:
                    if member.name in encrypted_names:
                        parts.append(
                            archive_member_warning(
                                path,
                                member.name,
                                "password_required",
                            )
                        )
                        continue

                    try:
                        copy_zip_member_safely(archive, member, tmp_path)
                    except Exception as exc:
                        parts.append(archive_member_warning(path, member.name, str(exc)))

            member_paths = extracted_member_paths(tmp_path, members, path)
            parts.append(extract_archive_member_texts(path, member_paths, archive_depth))
    except Exception as exc:
        return archive_warning(path, str(exc))

    return clean_text("\n".join(parts))


def seven_zip_item_is_dir(item) -> bool:
    is_directory = getattr(item, "is_directory", None)
    if callable(is_directory):
        return bool(is_directory())
    return bool(is_directory)


def archive_needs_password(archive) -> bool:
    needs_password = getattr(archive, "needs_password", None)
    if callable(needs_password):
        return bool(needs_password())
    return bool(needs_password)


def extract_7z(path: Path, archive_depth: int = 1) -> str:
    try:
        import py7zr
    except ImportError:
        return archive_warning(path, "py7zr_not_installed")

    try:
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp_path = Path(tmpdir)

            with py7zr.SevenZipFile(path, mode="r") as archive:
                if archive_needs_password(archive):
                    return archive_warning(path, "password_required")

                candidates = [
                    ArchiveMemberCandidate(
                        name=item.filename,
                        size=getattr(item, "uncompressed", 0),
                        is_dir=seven_zip_item_is_dir(item),
                    )
                    for item in archive.list()
                ]
                members = select_safe_archive_members(
                    candidates,
                    archive_path=path,
                    logger=logger,
                )
                if not members:
                    return ""

                archive.extract(
                    path=tmp_path,
                    targets=[member.name for member in members],
                )

            member_paths = extracted_member_paths(tmp_path, members, path)
            return extract_archive_member_texts(path, member_paths, archive_depth)
    except Exception as exc:
        return archive_warning(path, str(exc))


def rar_member_is_dir(item) -> bool:
    isdir = getattr(item, "isdir", None)
    if callable(isdir):
        return bool(isdir())
    return bool(getattr(item, "is_dir", False))


def detect_rar_backend() -> tuple[str | None, str | None]:
    for tool in RAR_BACKEND_CANDIDATES:
        path = shutil.which(tool)
        if path:
            return tool, path

    return None, None


def check_rar_backend(archive_path: Path | None = None) -> bool:
    tool, tool_path = detect_rar_backend()
    print(f"rar_backend_tool={tool or '-'}")
    print(f"rar_backend_path={tool_path or '-'}")

    try:
        import rarfile
    except ImportError as exc:
        print(f"rarfile_import=failed reason={exc}")
        return False

    print("rarfile_import=ok")

    if archive_path is None:
        return tool is not None

    try:
        with rarfile.RarFile(archive_path) as archive:
            names = archive.namelist()
        print(f"rarfile_open=ok members={len(names)}")
        return True
    except Exception as exc:
        print(f"rarfile_open=failed reason={exc}")
        return False


def extract_rar(path: Path, archive_depth: int = 1) -> str:
    try:
        import rarfile
    except ImportError:
        return archive_warning(path, "rarfile_not_installed")

    try:
        tool, _tool_path = detect_rar_backend()
        if not tool:
            return archive_warning(path, "rar_backend_not_found")

        with tempfile.TemporaryDirectory() as tmpdir:
            tmp_path = Path(tmpdir)

            with rarfile.RarFile(path) as archive:
                if archive_needs_password(archive):
                    return archive_warning(path, "password_required")

                candidates = [
                    ArchiveMemberCandidate(
                        name=item.filename,
                        size=getattr(item, "file_size", 0),
                        is_dir=rar_member_is_dir(item),
                    )
                    for item in archive.infolist()
                ]
                members = select_safe_archive_members(
                    candidates,
                    archive_path=path,
                    logger=logger,
                )

                for member in members:
                    try:
                        archive.extract(member.name, path=tmp_path)
                    except Exception as exc:
                        logger.warning(
                            "RAR member extraction failed: archive=%s member=%r reason=%s",
                            path,
                            member.name,
                            exc,
                        )

            member_paths = extracted_member_paths(tmp_path, members, path)
            return extract_archive_member_texts(path, member_paths, archive_depth)
    except Exception as exc:
        return archive_warning(path, str(exc))


def detect_zip_office_format(path: Path) -> str | None:
    try:
        if not zipfile.is_zipfile(path):
            return None
        with zipfile.ZipFile(path) as archive:
            names = {name.replace("\\", "/") for name in archive.namelist()}
    except (OSError, zipfile.BadZipFile):
        return None

    if "word/document.xml" in names:
        return "docx"
    if "xl/workbook.xml" in names:
        return "xlsx"

    return None


def extract_file(path: Path, archive_depth: int = 0) -> str:
    suffix = path.suffix.lower()

    if suffix not in {".zip", ".7z", ".rar"}:
        office_format = detect_zip_office_format(path)
        if office_format == "docx":
            return extract_docx(path)
        if office_format == "xlsx":
            return extract_xlsx(path)

    if suffix == ".docx":
        return extract_docx(path)

    if suffix == ".xlsx":
        return extract_xlsx(path)

    if suffix == ".xls":
        return extract_xls(path)

    if suffix == ".doc":
        return extract_doc_legacy(path)

    if suffix == ".pdf":
        return extract_pdf(path)

    if suffix == ".zip":
        if archive_depth >= MAX_ARCHIVE_DEPTH:
            return archive_warning(path, "archive_depth_limit")
        return extract_zip(path, archive_depth + 1)

    if suffix == ".7z":
        if archive_depth >= MAX_ARCHIVE_DEPTH:
            return archive_warning(path, "archive_depth_limit")
        return extract_7z(path, archive_depth + 1)

    if suffix == ".rar":
        if archive_depth >= MAX_ARCHIVE_DEPTH:
            return archive_warning(path, "archive_depth_limit")
        return extract_rar(path, archive_depth + 1)

    return ""


def empty_extraction_reason(path: Path) -> str:
    suffix = path.suffix.lower()

    remembered = remembered_empty_extraction_reason(path)
    if remembered:
        return remembered

    if not path.exists():
        return "file_not_found"

    if suffix == ".doc" and not shutil.which(LEGACY_DOC_TOOL):
        return f"{LEGACY_DOC_TOOL}_not_found"

    if suffix == ".rar":
        tool, _ = detect_rar_backend()
        if not tool:
            return "rar_backend_not_found"

    if suffix == ".7z":
        try:
            import py7zr  # noqa: F401
        except ImportError:
            return "py7zr_not_installed"

    supported_suffixes = {
        ".docx",
        ".xlsx",
        ".xls",
        ".doc",
        ".pdf",
        ".zip",
        ".7z",
        ".rar",
    }
    if suffix not in supported_suffixes:
        office_format = detect_zip_office_format(path)
        if office_format:
            return f"{office_format}_no_text_extracted"
        return f"unsupported_extension={suffix or 'none'}"

    if suffix == ".doc":
        return f"{LEGACY_DOC_TOOL}_no_text_extracted"

    if suffix == ".docx":
        return "docx_no_text_extracted"

    if suffix == ".xlsx":
        return "xlsx_no_cell_text"

    return "no_text_extracted"


def get_documents(limit: int, overwrite: bool) -> list[dict]:
    where = ""
    if not overwrite:
        where = "WHERE extracted_text IS NULL OR extracted_text = ''"

    query = f"""
        SELECT id, filename, storage_path, mime_type
        FROM documents
        {where}
        ORDER BY created_at DESC
        LIMIT %s;
    """

    with psycopg.connect(settings.database_url, row_factory=dict_row) as conn:
        with conn.cursor(row_factory=dict_row) as cur:
            cur.execute(query, (limit,))
            return cur.fetchall()


def save_extracted_text(document_id: str, text: str) -> None:
    with psycopg.connect(settings.database_url) as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                UPDATE documents
                SET extracted_text = %s
                WHERE id = %s;
                """,
                (text, document_id),
            )
        conn.commit()


def main() -> None:
    logging.basicConfig(
        level=logging.INFO,
        format="%(levelname)s:%(name)s:%(message)s",
    )

    parser = argparse.ArgumentParser(description="Extract text from downloaded documents")
    parser.add_argument("--limit", type=int, default=50)
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument(
        "--check-rar-backend",
        nargs="?",
        const="",
        default=None,
        metavar="ARCHIVE_PATH",
        help="Print visible rarfile backend and optionally try opening a RAR archive",
    )
    args = parser.parse_args()

    if args.check_rar_backend is not None:
        archive_path = Path(args.check_rar_backend) if args.check_rar_backend else None
        raise SystemExit(0 if check_rar_backend(archive_path) else 1)

    docs = get_documents(limit=args.limit, overwrite=args.overwrite)
    print(f"Documents to extract: {len(docs)}")

    extracted_count = 0

    for doc in docs:
        path = APP_DIR / doc["storage_path"]

        print()
        print(f"Document: {doc['filename']}")
        print(f"Path: {path}")

        if not path.exists():
            print("File not found")
            continue

        try:
            text = extract_file(path)
            text = clean_text(text)

            if not text:
                print("No text extracted")
                save_extracted_text(str(doc["id"]), "")
                continue

            save_extracted_text(str(doc["id"]), text)
            extracted_count += 1

            print(f"Extracted chars: {len(text)}")
            print(f"Preview: {text[:300].replace(chr(10), ' ')}")

        except Exception as exc:
            print(f"Failed: {exc}")

    print()
    print(f"Extracted documents: {extracted_count}")


if __name__ == "__main__":
    main()
