import json
import re
from dataclasses import dataclass
from datetime import date, datetime, time, timezone
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app import digest
from app.lead_notes import fetch_latest_lead_notes_for_tenders


EXCEL_EXPORT_TODAY = "today"
EXCEL_EXPORT_ALL = "all"
EXCEL_EXPORT_IN_WORK = "in_work"
EXCEL_EXPORT_BITRIX_IN_WORK = "bitrix_in_work"

EXCEL_EXPORT_SPECS = {
    EXCEL_EXPORT_TODAY: ("Новые клиентские лиды", "today"),
    EXCEL_EXPORT_ALL: ("Необработанные клиентские лиды", "unprocessed_leads"),
    EXCEL_EXPORT_IN_WORK: ("Лиды в работе", "in_work"),
    EXCEL_EXPORT_BITRIX_IN_WORK: ("Bitrix лиды", "bitrix_in_work"),
}

EXCEL_CELL_TEXT_LIMIT = 32000
ILLEGAL_CELL_CHAR_RE = re.compile(r"[\x00-\x08\x0B-\x0C\x0E-\x1F\x7F-\x9F\uD800-\uDFFF\uFFFE\uFFFF]")
FORMULA_PREFIXES = ("=", "+", "-", "@")

EXCEL_EXPORT_COLUMNS = [
    "№",
    "ID закупки / external_id",
    "Название",
    "Заказчик",
    "НМЦК",
    "Срок подачи",
    "Приоритет",
    "Уверенность",
    "Документы",
    "Рекомендация",
    "Сигнал по заказчику",
    "Гипотеза / история заказчика",
    "Возможные потребности",
    "Кого искать",
    "Контакт закупки",
    "Первая фраза",
    "Что выяснить",
    "Следующее действие",
    "Ссылка",
    "Статус",
    "Примечание менеджера",
]

WRAPPED_COLUMNS = {
    "Название",
    "Сигнал по заказчику",
    "Гипотеза / история заказчика",
    "Возможные потребности",
    "Кого искать",
    "Контакт закупки",
    "Первая фраза",
    "Что выяснить",
    "Следующее действие",
    "Примечание менеджера",
}
COLUMN_WIDTHS = {
    "№": 6,
    "ID закупки / external_id": 22,
    "Название": 42,
    "Заказчик": 32,
    "НМЦК": 16,
    "Срок подачи": 20,
    "Приоритет": 14,
    "Уверенность": 14,
    "Документы": 34,
    "Рекомендация": 28,
    "Сигнал по заказчику": 42,
    "Гипотеза / история заказчика": 42,
    "Возможные потребности": 38,
    "Кого искать": 34,
    "Контакт закупки": 34,
    "Первая фраза": 48,
    "Что выяснить": 42,
    "Следующее действие": 42,
    "Ссылка": 28,
    "Статус": 18,
    "Примечание менеджера": 42,
}

BITRIX_STATUS_NEW = "NEW"
BITRIX_SOURCE_ZAKUPKI360 = "1"
BITRIX_SOURCE_DESCRIPTION = "AI Tender Radar"
BITRIX_CURRENCY_RUB = "RUB"
BITRIX_OPENED_YES = "Y"
BITRIX_PLATFORM_TYPE_44_223 = "143"
BITRIX_PLATFORM_TYPE_615 = "144"
BITRIX_PLATFORM_TYPE_COMMERCIAL = "145"

BITRIX_EXPORT_COLUMNS = [
    "TITLE",
    "COMPANY_TITLE",
    "NAME",
    "PHONE",
    "EMAIL",
    "WEB",
    "COMMENTS",
    "STATUS_ID",
    "SOURCE_ID",
    "SOURCE_DESCRIPTION",
    "OPPORTUNITY",
    "CURRENCY_ID",
    "OPENED",
    "UF_CRM_1746003429",
    "UF_CRM_1746003445",
    "UF_CRM_1746003577",
    "UF_CRM_1746003587",
    "UF_CRM_1746003610",
    "UF_CRM_1746003622",
    "UF_CRM_1746003632",
    "UF_CRM_1746003677",
    "UF_CRM_1746004272",
    "UF_CRM_1746004281",
    "UF_CRM_1746003456",
    "UF_CRM_1746003471",
    "UF_CRM_1746003491",
    "UF_CRM_1746004349",
    "UF_CRM_1746004356",
    "UF_CRM_1746004365",
    "UF_CRM_1746004370",
    "UF_CRM_1746004376",
]

BITRIX_WRAPPED_COLUMNS = {"COMMENTS"}
BITRIX_PRESERVE_LEADING_PLUS_COLUMNS = {"PHONE"}
BITRIX_COLUMN_WIDTHS = {
    "TITLE": 52,
    "COMPANY_TITLE": 34,
    "NAME": 28,
    "PHONE": 20,
    "EMAIL": 28,
    "WEB": 32,
    "COMMENTS": 72,
    "STATUS_ID": 14,
    "SOURCE_ID": 12,
    "SOURCE_DESCRIPTION": 24,
    "OPPORTUNITY": 16,
    "CURRENCY_ID": 12,
    "OPENED": 10,
    "UF_CRM_1746003429": 20,
    "UF_CRM_1746003445": 32,
    "UF_CRM_1746003577": 18,
    "UF_CRM_1746003587": 22,
    "UF_CRM_1746003610": 28,
    "UF_CRM_1746003622": 20,
    "UF_CRM_1746003632": 20,
    "UF_CRM_1746003677": 18,
    "UF_CRM_1746004272": 28,
    "UF_CRM_1746004281": 34,
    "UF_CRM_1746003456": 32,
    "UF_CRM_1746003471": 32,
    "UF_CRM_1746003491": 32,
    "UF_CRM_1746004349": 22,
    "UF_CRM_1746004356": 36,
    "UF_CRM_1746004365": 18,
    "UF_CRM_1746004370": 18,
    "UF_CRM_1746004376": 36,
}


@dataclass(frozen=True)
class ExcelExportFile:
    filename: str
    content: bytes
    rows_count: int


class ExcelExportNoRows(ValueError):
    pass


def _safe_cell_value(value: Any, *, formula_prefixes: tuple[str, ...] = FORMULA_PREFIXES) -> Any:
    if value is None:
        return ""
    if isinstance(value, (list, tuple, set)):
        return _safe_cell_text(
            "\n".join(_plain_list_item(item) for item in value),
            formula_prefixes=formula_prefixes,
        )
    if isinstance(value, (bool, int, float)):
        return value
    if not isinstance(value, str):
        return _safe_cell_text(str(value), formula_prefixes=formula_prefixes)
    return _safe_cell_text(value, formula_prefixes=formula_prefixes)


def _plain_list_item(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, tuple, set)):
        return "\n".join(_plain_list_item(item) for item in value)
    return str(value)


def _safe_cell_text(value: str, *, formula_prefixes: tuple[str, ...] = FORMULA_PREFIXES) -> str:
    text = value.replace("\r\n", "\n").replace("\r", "\n")
    text = ILLEGAL_CELL_CHAR_RE.sub("", text)
    if text.startswith(formula_prefixes):
        text = f"'{text}"
    return text[:EXCEL_CELL_TEXT_LIMIT]


def _first_value(*values: Any) -> Any:
    for value in values:
        if value is None:
            continue
        if isinstance(value, str) and not value.strip():
            continue
        return value
    return None


def _as_dict(value: Any) -> dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _as_list(value: Any) -> list[Any]:
    if value is None:
        return []
    if isinstance(value, list):
        return value
    if isinstance(value, tuple):
        return list(value)
    return [value]


def _short_plain_text(value: Any, limit: int = 700) -> str:
    text = " ".join(str(value or "").split())
    if len(text) <= limit:
        return text
    return text[: limit - 1].rstrip() + "…"


def _non_negative_int(value: Any) -> int:
    try:
        parsed = int(float(str(value).strip()))
    except (TypeError, ValueError):
        return 0
    return max(0, parsed)


def _datetime_for_delta(value: Any, *, now: datetime) -> datetime | None:
    if not value:
        return None
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, time.max)
    elif isinstance(value, str):
        try:
            parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
        except ValueError:
            return None
    else:
        return None

    if parsed.tzinfo is None and now.tzinfo is not None:
        parsed = parsed.replace(tzinfo=now.tzinfo)
    elif parsed.tzinfo is not None and now.tzinfo is None:
        parsed = parsed.replace(tzinfo=None)
    return parsed


def _days_until_text(value: Any, *, now: datetime) -> str:
    deadline = _datetime_for_delta(value, now=now)
    if deadline is None:
        return "не указано"

    delta = deadline - now
    total_seconds = int(delta.total_seconds())
    if total_seconds < 0:
        return "просрочено"

    hours = total_seconds // 3600
    if hours < 48:
        return f"{hours} ч."

    days = hours // 24
    return f"{days} дн."


def _human_date(value: Any) -> str:
    if not value:
        return ""
    return digest.date_text(value)


def _human_money(value: Any) -> str:
    return digest.money(value)


def _bitrix_money_value(value: Any) -> Any:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, (int, float)):
        return value

    text = str(value).strip()
    if not text:
        return ""

    normalized = re.sub(r"[^\d,.\-]", "", text)
    if "," in normalized and "." in normalized:
        normalized = normalized.replace(",", "")
    elif "," in normalized:
        normalized = normalized.replace(",", ".")

    try:
        number = float(normalized)
    except ValueError:
        return ""
    if number.is_integer():
        return int(number)
    return number


def _bitrix_date_text(value: Any) -> str:
    if not value:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value.replace("Z", "+00:00")).date().isoformat()
        except ValueError:
            return _human_date(value) or value
    return _human_date(value)


def _bitrix_datetime_text(value: Any) -> str:
    if not value:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return datetime.combine(value, time.min).strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value.replace("Z", "+00:00")).strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            return _human_date(value) or value
    return _human_date(value)


def _money_for_phrase(value: Any) -> str:
    text = _human_money(value)
    return text if text else "не указана"


def _manual_document_reason(tender: dict[str, Any]) -> str | None:
    return (
        tender.get(digest.MANUAL_DOCUMENT_REASON_KEY)
        or tender.get(digest.TODAY_MANUAL_DOCUMENT_REASON_KEY)
        or digest.manual_document_reason_for_tender(tender)
    )


def _queue_status(tender: dict[str, Any], manual_reason: str | None) -> str:
    if manual_reason:
        return "document_needed"

    latest_status = digest.latest_sales_workflow_status(tender)
    if latest_status == digest.SALES_IN_WORK_STATUS:
        return "in_work"
    if latest_status == digest.SALES_RETURNED_STATUS:
        return "returned"
    if latest_status:
        return "unknown"
    return "new"


def _readiness(tender: dict[str, Any], manual_reason: str | None) -> str:
    if manual_reason == "preparation_blocked=marketplace_auth":
        return "marketplace_auth"
    if manual_reason == "preparation_no_valid_documents":
        return "no_valid_documents"
    if manual_reason == "preparation_exhausted=documents_missing":
        return "documents_missing"
    if manual_reason:
        return "need_documents"

    if digest.tender_has_extracted_document_text(tender):
        return "documents_ready_with_text"

    docs_count = _non_negative_int(tender.get("docs_count"))
    docs_with_text = _non_negative_int(tender.get("docs_with_text"))
    if docs_count <= 0:
        return "need_documents"
    if docs_with_text <= 0:
        return "no_valid_documents"
    return "unknown"


def _report_from_result(result: Any) -> dict[str, Any]:
    if not isinstance(result, dict):
        return {}
    report = result.get("report")
    return report if isinstance(report, dict) else {}


def _llm_status(tender: dict[str, Any], readiness: str) -> str:
    result = tender.get("llm_report_result")
    if isinstance(result, dict):
        status = str(result.get("status") or "").lower()
        if status == "failed" or result.get("error"):
            return "failed"

        report = _report_from_result(result)
        if report:
            if digest.llm_report_is_stale_without_document_text(tender, result):
                return "not_ready_for_llm"
            return "report_ready"

    if readiness in {
        "need_documents",
        "marketplace_auth",
        "no_valid_documents",
        "documents_missing",
    }:
        return "not_ready_for_llm"

    return "not_run"


def _llm_recommendation(tender: dict[str, Any]) -> str:
    direct = tender.get("llm_report_recommendation")
    if direct:
        return str(direct)

    report = _report_from_result(tender.get("llm_report_result"))
    recommendation = _as_dict(report.get("recommendation"))
    decision = recommendation.get("decision") or report.get("verdict")
    return str(decision or "")


def _procurement_number(tender: dict[str, Any], full: dict[str, Any]) -> str:
    value = _first_value(
        tender.get("procurement_number"),
        tender.get("reg_number"),
        tender.get("registration_number"),
        full.get("orderNumber"),
        full.get("purchaseNumber"),
        full.get("registrationNumber"),
        full.get("regNumber"),
        full.get("notificationNumber"),
        full.get("noticeNumber"),
    )
    return str(value or tender.get("external_id") or "")


def _price_value(tender: dict[str, Any], full: dict[str, Any]) -> Any:
    return _first_value(
        tender.get("initial_price"),
        tender.get("max_price"),
        tender.get("price"),
        full.get("maxPrice"),
        full.get("nmck"),
        full.get("initialPrice"),
    )


def _published_at(tender: dict[str, Any], full: dict[str, Any]) -> Any:
    return _first_value(
        tender.get("published_at"),
        tender.get("publication_date"),
        tender.get("publish_date"),
        full.get("publishDate"),
        full.get("publicationDate"),
        full.get("createDate"),
    )


def _deadline_at(tender: dict[str, Any], full: dict[str, Any]) -> Any:
    return _first_value(
        tender.get("deadline_at"),
        tender.get("application_deadline"),
        tender.get("end_date"),
        tender.get("bidding_deadline"),
        full.get("endDate"),
        full.get("filingEndDate"),
        full.get("applicationDeadline"),
    )


def _contact_fields(full: dict[str, Any]) -> tuple[str, str, str]:
    contact = _as_dict(full.get("contact") or full.get("contactInfo"))
    person = _first_value(
        full.get("contactPerson"),
        full.get("contactName"),
        contact.get("person"),
        contact.get("name"),
    )
    phone = _first_value(
        full.get("contactPhone"),
        full.get("phone"),
        contact.get("phone"),
        contact.get("telephone"),
    )
    email = _first_value(
        full.get("contactEMail"),
        full.get("contactEmail"),
        full.get("email"),
        contact.get("email"),
    )
    return str(person or ""), str(phone or ""), str(email or "")


def _why_interesting(tender: dict[str, Any]) -> str:
    result = _as_dict(tender.get("result") or tender.get("rule_based_result"))
    candidates: list[str] = []
    for key in ("why_relevant", "positive_factors", "manual_checks"):
        for item in _as_list(result.get(key)):
            text = _short_plain_text(item, limit=240)
            if text and text not in candidates:
                candidates.append(text)
            if len(candidates) >= 3:
                return "\n".join(candidates)
    return ""


def _business_action(tender: dict[str, Any], recommendation: str, manual_reason: str | None) -> str:
    if manual_reason:
        return digest.manual_document_action_text(manual_reason)

    try:
        block = digest.digest_business_assessment_block(tender)
    except Exception:
        block = ""

    for line in block.splitlines():
        if line.startswith("Бизнес-действие:"):
            return line.split(":", 1)[1].strip()

    return digest.next_action_block(recommendation, tender)


def _contact_reason(
    tender: dict[str, Any],
    full: dict[str, Any],
    price: Any,
    deadline: Any,
    manual_reason: str | None,
) -> str:
    if manual_reason:
        return "Повод: запросить/скачать КД/ТЗ/спецификацию и выйти на ответственного."

    subject = _first_value(
        full.get("categoryName"),
        full.get("subjectName"),
        full.get("objectName"),
        tender.get("title"),
    )
    return (
        f"Закупка по {_short_plain_text(subject, limit=120)}, "
        f"сумма {_money_for_phrase(price)}, срок {digest.date_text(deadline)}. "
        "Повод: уточнить задачу и выйти на технического ответственного."
    )


def _feedback_comment(value: Any) -> str:
    if not value:
        return ""

    payload = value
    if isinstance(value, str):
        try:
            payload = json.loads(value)
        except json.JSONDecodeError:
            return _short_plain_text(value)

    if isinstance(payload, dict):
        for key in ("message", "comment", "text", "note"):
            if payload.get(key):
                return _short_plain_text(payload.get(key))
        return ""

    return _short_plain_text(payload)


def _lead_recommendation(tender: dict[str, Any]) -> str:
    queue_decision = digest.operational_queue_decision_from_row(tender)
    if queue_decision:
        return queue_decision
    direct = digest.lead_report_stored_recommendation_from_row(tender)
    if direct:
        return direct
    normalized = digest.normalized_recommendation(
        tender.get("llm_report_recommendation") or tender.get("recommendation")
    )
    return normalized or "go"


def _lead_status_text(tender: dict[str, Any]) -> str:
    status = digest.latest_sales_workflow_status(tender)
    return {
        digest.SALES_IN_WORK_STATUS: "В работе",
    }.get(str(status or ""), "Новый")


def _list_cell(items: list[str]) -> str:
    return "\n".join(items)


def _nested_first(mapping: dict[str, Any], *keys: str) -> Any:
    for key in keys:
        value = mapping.get(key)
        if value is not None and not (isinstance(value, str) and not value.strip()):
            return value
    return None


def _organizer_dict(full: dict[str, Any]) -> dict[str, Any]:
    for key in ("organizer", "organizerInfo", "placer", "placerInfo"):
        value = full.get(key)
        if isinstance(value, dict):
            return value
    return {}


def _bitrix_title(tender: dict[str, Any], full: dict[str, Any]) -> str:
    customer = digest.get_customer_name(tender, full)
    title = _short_plain_text(tender.get("title"), limit=170)
    procurement_number = _procurement_number(tender, full)
    parts = [part for part in (customer, title, procurement_number) if part]
    return _short_plain_text(" / ".join(parts), limit=250)


def _bitrix_platform_type_id(tender: dict[str, Any], full: dict[str, Any], url: str) -> str:
    law_text = " ".join(
        str(value or "")
        for value in (
            tender.get("law"),
            full.get("law"),
            full.get("lawId"),
            full.get("lawName"),
            full.get("fz"),
            full.get("typeName"),
            full.get("purchaseType"),
        )
    ).lower()
    url_text = str(url or "").lower()
    if "615" in law_text:
        return BITRIX_PLATFORM_TYPE_615
    if any(marker in law_text for marker in ("коммер", "змо", "commercial")):
        return BITRIX_PLATFORM_TYPE_COMMERCIAL
    if any(marker in law_text for marker in ("44", "223")) or "notice223" in url_text:
        return BITRIX_PLATFORM_TYPE_44_223
    return ""


def _bitrix_comments(
    tender: dict[str, Any],
    full: dict[str, Any],
    report: dict[str, Any],
    *,
    price: Any,
    deadline: Any,
    url: str,
    manager_note: str,
) -> str:
    possible_needs = digest.lead_report_possible_needs(report)
    who_to_find = digest.lead_report_who_to_find_values(tender, report, full)
    questions = digest.lead_report_discovery_questions(tender, report)
    lines: list[str] = ["AI Tender Radar"]

    if manager_note.strip():
        lines.extend(("", "Примечание менеджера:", manager_note.strip()))

    summary_items = [
        ("ID закупки", tender.get("external_id") or tender.get("tender_id") or tender.get("id")),
        ("Номер закупки", _procurement_number(tender, full)),
        ("Заказчик", digest.get_customer_name(tender, full)),
        ("НМЦК", _human_money(price)),
        ("Срок подачи", _human_date(deadline)),
        ("Ссылка", url),
        ("Приоритет", digest.lead_report_priority_text(digest.lead_report_priority_from_row(tender))),
        ("Уверенность", digest.lead_report_confidence_from_row(tender, report)),
        ("Документы", digest.lead_report_document_status_label(tender, report)),
    ]
    lines.append("")
    for label, value in summary_items:
        if value:
            lines.append(f"{label}: {value}")

    text_blocks = [
        ("Сигнал по заказчику", digest.lead_report_signal_text(tender, report)),
        ("Гипотеза / история заказчика", digest.lead_report_customer_story_text(report)),
        ("Первая фраза", digest.lead_report_opening_phrase(tender, report)),
        ("Следующее действие", digest.lead_report_next_action(report)),
        ("Контакт закупки", digest.procurement_contact_text(full) or ""),
    ]
    for label, value in text_blocks:
        if value:
            lines.extend(("", f"{label}:", str(value)))

    list_blocks = [
        ("Возможные потребности", possible_needs),
        ("Кого искать", who_to_find),
        ("Что выяснить", questions),
    ]
    for label, values in list_blocks:
        cleaned = [str(value).strip() for value in values if str(value).strip()]
        if cleaned:
            lines.extend(("", f"{label}:"))
            lines.extend(f"- {value}" for value in cleaned)

    return "\n".join(lines).strip()


def bitrix_row_from_tender(
    tender: dict[str, Any],
    *,
    now: datetime | None = None,
    index: int | None = None,
) -> dict[str, Any]:
    del now, index
    full = digest.raw_full(tender)
    report = digest.lead_report_body(tender)
    price = _price_value(tender, full)
    deadline = _deadline_at(tender, full)
    published = _published_at(tender, full)
    url = digest.procedure_url(tender, full) or ""
    contact_person, contact_phone, contact_email = _contact_fields(full)
    manager_note = str(tender.get("manager_note") or "")
    organizer = _organizer_dict(full)

    return {
        "TITLE": _bitrix_title(tender, full),
        "COMPANY_TITLE": digest.get_customer_name(tender, full),
        "NAME": contact_person,
        "PHONE": contact_phone,
        "EMAIL": contact_email,
        "WEB": url,
        "COMMENTS": _bitrix_comments(
            tender,
            full,
            report,
            price=price,
            deadline=deadline,
            url=url,
            manager_note=manager_note,
        ),
        "STATUS_ID": BITRIX_STATUS_NEW,
        "SOURCE_ID": BITRIX_SOURCE_ZAKUPKI360,
        "SOURCE_DESCRIPTION": BITRIX_SOURCE_DESCRIPTION,
        "OPPORTUNITY": _bitrix_money_value(price),
        "CURRENCY_ID": BITRIX_CURRENCY_RUB,
        "OPENED": BITRIX_OPENED_YES,
        "UF_CRM_1746003429": str(
            tender.get("external_id")
            or tender.get("tender_id")
            or tender.get("id")
            or _procurement_number(tender, full)
        ),
        "UF_CRM_1746003445": url,
        "UF_CRM_1746003577": _bitrix_date_text(published),
        "UF_CRM_1746003587": _bitrix_datetime_text(deadline),
        "UF_CRM_1746003610": str(
            _first_value(
                tender.get("procedure_type"),
                full.get("purchaseMethodName"),
                full.get("placingWayName"),
                full.get("methodName"),
            )
            or ""
        ),
        "UF_CRM_1746003622": str(
            _first_value(
                full.get("statusName"),
                full.get("stageName"),
                full.get("stateName"),
                tender.get("status"),
            )
            or ""
        ),
        "UF_CRM_1746003632": _procurement_number(tender, full),
        "UF_CRM_1746003677": _bitrix_platform_type_id(tender, full, url),
        "UF_CRM_1746004272": str(
            _first_value(
                full.get("platformName"),
                full.get("etpName"),
                full.get("marketplaceName"),
                full.get("tradePlatformName"),
            )
            or ""
        ),
        "UF_CRM_1746004281": str(
            _first_value(
                full.get("deliveryPlace"),
                full.get("deliveryAddress"),
                full.get("placeOfDelivery"),
                tender.get("region"),
            )
            or ""
        ),
        "UF_CRM_1746003456": str(
            _first_value(
                tender.get("z360_url"),
                full.get("personalAccountUrl"),
                full.get("privateUrl"),
            )
            or ""
        ),
        "UF_CRM_1746003471": str(
            _first_value(full.get("customerUrl"), full.get("customerLink"), full.get("customerCardUrl")) or ""
        ),
        "UF_CRM_1746003491": str(
            _first_value(full.get("organizerUrl"), full.get("organizerLink"), full.get("placerUrl")) or ""
        ),
        "UF_CRM_1746004349": str(
            _first_value(_nested_first(organizer, "id", "z360Id"), full.get("organizerId"), full.get("placerId"))
            or ""
        ),
        "UF_CRM_1746004356": str(
            _first_value(
                _nested_first(organizer, "fullName", "name"),
                full.get("organizerFullName"),
                full.get("organizerName"),
                full.get("placerFullName"),
                full.get("placerName"),
            )
            or ""
        ),
        "UF_CRM_1746004365": str(
            _first_value(_nested_first(organizer, "inn", "INN"), full.get("organizerInn"), full.get("placerInn"))
            or ""
        ),
        "UF_CRM_1746004370": str(
            _first_value(_nested_first(organizer, "kpp", "KPP"), full.get("organizerKpp"), full.get("placerKpp"))
            or ""
        ),
        "UF_CRM_1746004376": str(
            _first_value(
                _nested_first(organizer, "address", "legalAddress"),
                full.get("organizerAddress"),
                full.get("placerAddress"),
            )
            or ""
        ),
    }


def excel_row_from_tender(
    tender: dict[str, Any],
    *,
    now: datetime | None = None,
    index: int | None = None,
) -> dict[str, Any]:
    del now
    full = digest.raw_full(tender)
    report = digest.lead_report_body(tender)
    recommendation = _lead_recommendation(tender)
    price = _price_value(tender, full)
    deadline = _deadline_at(tender, full)
    url = digest.procedure_url(tender, full) or ""
    possible_needs = digest.lead_report_possible_needs(report)
    who_to_find = digest.lead_report_who_to_find_values(tender, report, full)
    questions = digest.lead_report_discovery_questions(tender, report)

    return {
        "№": index or "",
        "ID закупки / external_id": str(
            tender.get("external_id")
            or tender.get("tender_id")
            or tender.get("id")
            or _procurement_number(tender, full)
        ),
        "Название": str(tender.get("title") or ""),
        "Заказчик": digest.get_customer_name(tender, full),
        "НМЦК": _human_money(price),
        "Срок подачи": _human_date(deadline),
        "Приоритет": digest.lead_report_priority_text(digest.lead_report_priority_from_row(tender)),
        "Уверенность": digest.lead_report_confidence_from_row(tender, report),
        "Документы": digest.lead_report_document_status_label(tender, report),
        "Рекомендация": digest.call_first_recommendation_text(recommendation),
        "Сигнал по заказчику": digest.lead_report_signal_text(tender, report),
        "Гипотеза / история заказчика": digest.lead_report_customer_story_text(report),
        "Возможные потребности": _list_cell(possible_needs),
        "Кого искать": _list_cell(who_to_find),
        "Контакт закупки": digest.procurement_contact_text(full) or "",
        "Первая фраза": digest.lead_report_opening_phrase(tender, report),
        "Что выяснить": _list_cell(questions),
        "Следующее действие": digest.lead_report_next_action(report),
        "Ссылка": url,
        "Статус": _lead_status_text(tender),
        "Примечание менеджера": str(tender.get("manager_note") or ""),
    }


def select_excel_export_source_rows(
    mode: str,
    *,
    today_limit: int = 100,
    in_work_limit: int = 100,
    now: datetime | None = None,
) -> list[dict[str, Any]]:
    def visible_queue_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
        return [
            dict(row)
            for row in rows
            if digest.operational_lead_hidden_reason(row) is None
        ]

    def visible_in_work_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
        return [
            dict(row)
            for row in rows
            if not digest.hidden_by_latest_lead_triage_reject(row)
        ]

    if mode == EXCEL_EXPORT_TODAY:
        rows, _ = digest.select_operational_lead_queue_rows(limit=today_limit, now=now)
        return visible_queue_rows(rows)

    if mode == EXCEL_EXPORT_ALL:
        rows, _ = digest.select_operational_lead_all_rows(limit=None, now=now)
        return visible_queue_rows(rows)

    if mode in {EXCEL_EXPORT_IN_WORK, EXCEL_EXPORT_BITRIX_IN_WORK}:
        rows, _ = digest.select_operational_lead_in_work_rows(
            limit=in_work_limit,
            now=now,
        )
        return visible_in_work_rows(rows)

    raise ValueError(f"Unknown Excel export mode: {mode}")


def build_excel_export_rows(
    mode: str,
    *,
    now: datetime | None = None,
) -> list[dict[str, Any]]:
    current_time = now or datetime.now(timezone.utc)
    source_rows = select_excel_export_source_rows(mode, now=current_time)
    row_builder = bitrix_row_from_tender if mode == EXCEL_EXPORT_BITRIX_IN_WORK else excel_row_from_tender
    notes_by_tender_id = fetch_latest_lead_notes_for_tenders(
        row.get("tender_id") or row.get("id")
        for row in source_rows
    )
    enriched_rows: list[dict[str, Any]] = []
    for row in source_rows:
        tender_id = str(row.get("tender_id") or row.get("id") or "")
        enriched = dict(row)
        if tender_id in notes_by_tender_id:
            enriched["manager_note"] = notes_by_tender_id[tender_id]
        enriched_rows.append(enriched)

    return [
        row_builder(row, now=current_time, index=index)
        for index, row in enumerate(enriched_rows, start=1)
    ]


def _export_columns_for_mode(mode: str) -> list[str]:
    if mode == EXCEL_EXPORT_BITRIX_IN_WORK:
        return BITRIX_EXPORT_COLUMNS
    return EXCEL_EXPORT_COLUMNS


def _wrapped_columns_for_mode(mode: str) -> set[str]:
    if mode == EXCEL_EXPORT_BITRIX_IN_WORK:
        return BITRIX_WRAPPED_COLUMNS
    return WRAPPED_COLUMNS


def _column_widths_for_mode(mode: str) -> dict[str, int]:
    if mode == EXCEL_EXPORT_BITRIX_IN_WORK:
        return BITRIX_COLUMN_WIDTHS
    return COLUMN_WIDTHS


def _preserve_leading_plus_columns_for_mode(mode: str) -> set[str]:
    if mode == EXCEL_EXPORT_BITRIX_IN_WORK:
        return BITRIX_PRESERVE_LEADING_PLUS_COLUMNS
    return set()


def _column_width(header: str, values: list[Any], *, column_widths: dict[str, int] | None = None) -> int:
    widths = column_widths or COLUMN_WIDTHS
    if header in widths:
        return widths[header]

    max_len = len(header)
    for value in values[:200]:
        max_len = max(max_len, max(len(line) for line in str(value or "").splitlines() or [""]))
    return max(10, min(max_len + 2, 56))


def build_excel_workbook_bytes(
    *,
    sheet_name: str,
    rows: list[dict[str, Any]],
    table_name: str | None = None,
    columns: list[str] | None = None,
    wrapped_columns: set[str] | None = None,
    column_widths: dict[str, int] | None = None,
    preserve_leading_plus_columns: set[str] | None = None,
) -> bytes:
    del table_name
    workbook_columns = columns or EXCEL_EXPORT_COLUMNS
    workbook_wrapped_columns = wrapped_columns or WRAPPED_COLUMNS
    workbook_preserve_plus = preserve_leading_plus_columns or set()
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name[:31]
    worksheet.freeze_panes = "A2"

    worksheet.append([_safe_cell_value(column) for column in workbook_columns])
    for row in rows:
        values = []
        for column in workbook_columns:
            formula_prefixes = FORMULA_PREFIXES
            if column in workbook_preserve_plus:
                formula_prefixes = tuple(prefix for prefix in FORMULA_PREFIXES if prefix != "+")
            values.append(_safe_cell_value(row.get(column, ""), formula_prefixes=formula_prefixes))
        worksheet.append(values)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for column_index, header in enumerate(workbook_columns, start=1):
        column_letter = get_column_letter(column_index)
        values = [
            worksheet.cell(row=row_number, column=column_index).value
            for row_number in range(2, worksheet.max_row + 1)
        ]
        worksheet.column_dimensions[column_letter].width = _column_width(
            header,
            values,
            column_widths=column_widths,
        )
        for row_number in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row_number, column=column_index)
            cell.alignment = Alignment(
                wrap_text=header in workbook_wrapped_columns,
                vertical="top",
            )

    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(workbook_columns))}{worksheet.max_row}"

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def build_excel_export_file(
    mode: str,
    *,
    now: datetime | None = None,
) -> ExcelExportFile:
    if mode not in EXCEL_EXPORT_SPECS:
        raise ValueError(f"Unknown Excel export mode: {mode}")

    current_time = now or datetime.now()
    sheet_name, filename_part = EXCEL_EXPORT_SPECS[mode]
    rows = build_excel_export_rows(mode, now=current_time)
    if not rows:
        raise ExcelExportNoRows("No tenders for Excel export")

    content = build_excel_workbook_bytes(
        sheet_name=sheet_name,
        rows=rows,
        columns=_export_columns_for_mode(mode),
        wrapped_columns=_wrapped_columns_for_mode(mode),
        column_widths=_column_widths_for_mode(mode),
        preserve_leading_plus_columns=_preserve_leading_plus_columns_for_mode(mode),
    )
    filename = f"ai_tender_radar_{filename_part}_{current_time:%Y%m%d_%H%M}.xlsx"
    return ExcelExportFile(filename=filename, content=content, rows_count=len(rows))
